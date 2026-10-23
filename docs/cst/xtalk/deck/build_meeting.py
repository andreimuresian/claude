"""Crosstalk from cross-section modes: slide deck in the company template (python-pptx).
Template = the company deck docs/CST_line_extraction.pptx with its slides removed."""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
HERE = os.path.dirname(os.path.abspath(__file__)); IMG = os.path.join(HERE, "..")
NAVY = RGBColor(0x0E, 0x22, 0x68); RED = RGBColor(0xF2, 0x00, 0x13); INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x60, 0x6E); CARD = RGBColor(0xF1, 0xF3, 0xF8); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW = RGBColor(0xF6, 0xF7, 0xFA); BOX = RGBColor(0xD9, 0xDC, 0xE3); GREEN = RGBColor(0x2E, 0x8B, 0x57)
FONT = "Arial"
prs = Presentation(os.path.join(HERE, "..", "..", "..", "CST_line_extraction.pptx"))
lst = prs.slides._sldIdLst
for s in list(lst):
    prs.part.drop_rel(s.get(qn("r:id"))); lst.remove(s)
L_TITLE, L_BLANK = prs.slide_layouts[11], prs.slide_layouts[12]

def runs(p, t, size, color, bold=False):
    for tok in re.split(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})", t):
        if not tok: continue
        b, base, s = bold, 0, tok
        if tok.startswith("**"): b, s = True, tok[2:-2]
        elif tok.startswith("_{"): base, s = -25000, tok[2:-1]
        elif tok.startswith("^{"): base, s = 30000, tok[2:-1]
        r = p.add_run(); r.text = s; f = r.font
        f.name, f.size, f.bold = FONT, Pt(size), b; f.color.rgb = color
        if base: r._r.get_or_add_rPr().set("baseline", str(base))

def text(s, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sa=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02); tf.vertical_anchor = anchor
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(sa)
        sz, col, bd = size, color, bold
        if isinstance(para, tuple):
            para, o = para; sz, col, bd = o.get("size", sz), o.get("color", col), o.get("bold", bd)
        runs(p, para, sz, col, bd)
    return tb

def title(s, t, sub=None):
    text(s, 0.55, 0.42, 10.6, 0.7, t, size=28)
    if sub: text(s, 0.6, 1.02, 11.8, 0.4, sub, size=15, color=NAVY, bold=True)

def rect(s, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh

def pic(s, name, x, y, w, h):
    path = os.path.join(IMG, name); W, H = Image.open(path).size; sc = min(w / W, h / H); ww, hh = W * sc, H * sc
    return s.shapes.add_picture(path, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))

def fig(s, name, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=BOX); pic(s, name, x + 0.05, y + 0.05, w - 0.1, h - 0.1)

def card(s, x, y, w, h, head, body, hc=NAVY, size=13.5):
    rect(s, x, y, w, h)
    text(s, x + 0.15, y + 0.08, w - 0.3, h - 0.12, [(head, {"bold": True, "color": hc, "size": size + 1.5})] + [(b, {"size": size}) for b in body], sa=3)

def stat(s, x, y, w, h, big, small, col=NAVY):
    rect(s, x, y, w, h)
    text(s, x, y + 0.15, w, 1.0, big, size=40, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, y + 1.2, w - 0.4, h - 1.3, small, size=13, color=INK, align=PP_ALIGN.CENTER)

def table(s, x, y, w, rows, col_w, size=11, row_h=0.42, hl=()):
    gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))); tb = gt.table
    pr = tb._tbl.tblPr
    for k in ("bandRow", "firstRow"): pr.set(k, "0")
    st = pr.find(qn("a:tableStyleId"))
    if st is not None: pr.remove(st)
    tot = sum(col_w)
    for j, cw in enumerate(col_w): tb.columns[j].width = Inches(w * cw / tot)
    for i, r in enumerate(rows):
        tb.rows[i].height = Inches(row_h)
        for j, v in enumerate(r):
            c = tb.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            runs(p, str(v), size, WHITE if i == 0 else INK, bold=(i == 0 or j == 0))
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xFD, 0xEC, 0xEE) if i in hl else (ROW if i % 2 else WHITE))

def arrow(s, x, y, w, h=0.4):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background(); a.shadow.inherit = False

new = lambda: prs.slides.add_slide(L_BLANK)
MI = os.path.join(HERE, "meeting_img"); RI = os.path.join(HERE, "..", "report", "img")
def figp(s, path, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=BOX)
    W, H = Image.open(path).size; sc = min((w - .1) / W, (h - .1) / H); ww, hh = W * sc, H * sc
    s.shapes.add_picture(path, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))
def eq(s, name, x, y, h=None, w=None):
    path = os.path.join(MI, name + ".png"); W, H = Image.open(path).size
    if h is None: h = w * H / W
    if w is None: w = h * W / H
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h)); return w, h
def box(s, x, y, w, h, head, sub, fill=CARD, hc=NAVY):
    rect(s, x, y, w, h, fill=fill, line=BOX)
    text(s, x + .1, y + .08, w - .2, h - .1, [(head, {"bold": True, "color": hc, "size": 14})] + [(t, {"size": 13, "color": INK}) for t in sub], sa=2)
def varrow(s, x, y, h, w=0.35):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background(); a.shadow.inherit = False

# 1 title
s = prs.slides.add_slide(L_TITLE)
text(s, 0, 2.2, 13.33, 1.4, ["Electrical and electro-optic crosstalk:", "from CST to INTERCONNECT"], size=36, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0, 3.75, 13.33, 0.8, "Two adjacent TFLN GSG electrodes · uniform line, pitch 200 µm · Python twin, GUI and INTERCONNECT model", size=18, color=MUTED, align=PP_ALIGN.CENTER)

# 2 summary
s = new(); title(s, "Summary", "A crosstalk model for any electrode length, in the GUI and in INTERCONNECT")
stat(s, 0.5, 1.6, 3.95, 2.3, "−47 dB", "FEXT at 30 GHz, 14 mm (−43 dB at 60 GHz)")
stat(s, 4.69, 1.6, 3.95, 2.3, "−47 dB", "EO crosstalk B/A at 30 GHz, 14 mm (−36 dB at 60 GHz)", col=RED)
stat(s, 8.88, 1.6, 3.95, 2.3, "10^{−14}", "INTERCONNECT element vs Python twin (code check)")
card(s, 0.5, 4.2, 6.05, 2.7, "Main message", ["Two coupled lines = an **even** line + an **odd** line, each the single-line element.",
     "Inputs: average line from 3-D CST; even/odd split from the 2-D and eigenmode solvers.",
     "Crosstalk vs length follows from closed-form formulas: no segmentation, no fitting."])
card(s, 6.78, 4.2, 6.05, 2.7, "Open points", ["Impedance split δ = 0.21–0.61 % (definition-dependent); 1 mm 3-D NEXT favours 0.23 %.",
     "Port / pad coupling is not in the model.", "Full MZM cascade: built, first INTERCONNECT run pending."], hc=RED)

# 3 workflow
s = new(); title(s, "Workflow", "Each simulation is used for what it resolves")
box(s, 0.5, 1.6, 3.9, 1.55, "3-D CST, time domain (coworker)", ["Pair, 1 mm and 5 mm, gold, real ports", "→ average line: loss, n, Z"])
box(s, 0.5, 3.35, 3.9, 1.55, "2-D CST port modes", ["One port over both lines, 6 modes", "→ Δn(f) of even/odd (PEC metal)"])
box(s, 0.5, 5.1, 3.9, 1.55, "CST eigenmode, one 10 µm cell", ["Gold vs PEC, even and odd", "→ gold correction of Δn (×1.146), δ band"])
for y in (2.2, 3.95, 5.7): arrow(s, 4.55, y, 0.75)
box(s, 5.45, 2.4, 3.4, 3.4, "Mode tables (even, odd)", ["γ_{e,o} = γ ± jΔβ/2", "Z_{e,o} = Z (1 ± δ/2)", "", "coupled_even.txt", "coupled_odd.txt", "", "valid for any length L"], fill=RGBColor(0xFD, 0xEC, 0xEE), hc=RED)
arrow(s, 9.0, 3.0, 0.75); arrow(s, 9.0, 4.8, 0.75)
box(s, 9.9, 2.25, 2.95, 1.6, "Python twin + GUI", ["coupled_line.py", "4 graphs, any length"])
box(s, 9.9, 4.35, 2.95, 1.6, "INTERCONNECT", ["\"Coupled TL\" scripted element", "full MZM cascade"])
text(s, 0.5, 6.75, 12.3, 0.4, "Why not crosstalk straight from 3-D: the 3-D runs' numerical error (|S_{21} − S_{43}| = −56 dB at 5 mm) is as large as the crosstalk itself.", size=12, color=MUTED)

# 4 two modes
s = new(); title(s, "Only the even and odd coplanar modes are needed", "A GSG launch excites line A in its coplanar mode = ½ even + ½ odd")
figp(s, os.path.join(RI, "modes.jpg"), 0.5, 1.55, 8.0, 3.6)
card(s, 8.75, 1.55, 4.1, 3.6, "Mode identification (not by eye)", ["7 conductors → exactly 6 quasi-TEM modes.",
     "Modes 5/6: V_{A} = ±V_{B} (from exported fields); each line's grounds equal, net current zero.",
     "Near-degenerate pair at the coplanar index (Δn = 1.5·10^{−3})."])
card(s, 0.5, 5.35, 6.05, 1.55, "Why modes 1–4 are not excited", ["Grounds tied + current returning on own grounds → zero amplitude for the slot and line-vs-frame modes."])
card(s, 6.78, 5.35, 6.07, 1.55, "Condition", ["Holds for a symmetric GSG launch. Asymmetric pads, bends without ground bridges or poor ports can excite them."], hc=RED)

# 5 inputs
s = new(); title(s, "Model inputs: index split Δn and impedance split δ", "Two independent numbers from the two couplings of the pair")
eq(s, "eq_coupling", 0.6, 1.6, h=0.62)
text(s, 0.6, 2.35, 6.2, 1.0, ["k_{L}, k_{C}: inductive and capacitive coupling. In a homogeneous dielectric k_{L} = k_{C} → Δn = 0 (no FEXT)."], size=13)
table(s, 0.5, 3.35, 7.6, [["", "Δn (30 GHz)", "δ = Z_{e}/Z_{o} − 1"],
    ["2-D port modes, PEC", "0.00151", "0.30–0.45 %"], ["Eigenmode, PEC", "0.00145", "+0.53 %"],
    ["Eigenmode, gold", "0.00173", "+0.21 % … +0.61 %"], ["Used in the model", "2-D × 1.146", "0.21–0.61 % (band)"]], [2.6, 1.6, 2.2], size=12, hl=(4,))
card(s, 8.35, 1.55, 4.5, 2.55, "How Δn is extracted", ["2-D: n = βc/(2πf) of modes 5, 6.", "Eigenmode: fixed β, Δn/n = (f_{o} − f_{e})/f.", "Gold raises Δn by 15–20 % vs PEC."])
card(s, 8.35, 4.25, 4.5, 2.65, "Why δ is a band", ["δ depends on the impedance definition: power–voltage +0.61 %, signal-to-own-grounds +0.21 %.",
     "Physical (non-TEM, lossy gold, ground currents), not mesh (refinement: 2 %)."], hc=RED)
text(s, 0.5, 5.75, 7.6, 1.1, ["Model: same loss for both modes; n and Z split around the 3-D line:"], size=13)
eq(s, "eq_modes", 0.6, 6.15, h=0.55)

# 6 model
s = new(); title(s, "The model: even line + odd line, recombined", "Exact for a uniform pair; every quantity is closed-form")
text(s, 0.6, 1.6, 12, 0.4, "**1.** Each mode is a single uniform line between z_{0} = 50 Ω ports:", size=15)
eq(s, "eq_single", 0.9, 2.05, h=0.75)
text(s, 0.6, 3.05, 12, 0.4, "**2.** Driving A1 alone = ½ even + ½ odd; the other line picks up the difference:", size=15)
eq(s, "eq_xt", 0.9, 3.5, h=0.62)
text(s, 0.6, 4.4, 12, 0.4, "**3.** Line voltages recombine the same way (used for the EO response):", size=15)
eq(s, "eq_VAB", 0.9, 4.85, h=0.55)
card(s, 0.5, 5.7, 12.35, 1.2, "Checks", ["Matches a brute-force solution of the coupled telegrapher equations to 10^{−13} (S) and 10^{−7} (modulation)."])

# 7 NEXT vs FEXT
s = new(); title(s, "NEXT is set by δ, FEXT by Δn", "First order in the even/odd differences (ΔZ = Zδ, Δγ = jΔβ)")
eq(s, "eq_next", 0.6, 1.65, h=0.62); eq(s, "eq_fext", 0.6, 2.5, h=0.62)
card(s, 0.5, 3.4, 4.1, 3.5, "Why", ["**NEXT** = reflection difference: ∂S_{11}/∂Z is first order → ∝ δ. Δn enters only through Γ (mismatch).",
     "**FEXT** = transmission difference: ∂S_{21}/∂t ≈ 1 → ∝ Δn. δ enters only as Γ·δ (second order).",
     "14 mm, 30 GHz: Δn term −42 dB; δ term ≤ −60 dB."])
figp(s, os.path.join(RI, "nulls.png"), 4.8, 3.4, 8.05, 3.5)
text(s, 7.4, 1.6, 5.4, 0.5, ["NEXT saturates at ≈ δ/2 and has nulls at"], size=14)
eq(s, "eq_nulls", 7.4, 2.05, h=0.62)
text(s, 7.4, 2.75, 5.4, 0.5, "≈ 5.5 GHz spacing at 14 mm", size=13, color=MUTED)

# 8 validation
s = new(); title(s, "Model vs 3-D CST", "1 mm agrees; 5 mm NEXT above 25 GHz is open (3-D at its error floor)")
figp(s, os.path.join(MI, "validation.png"), 0.5, 1.5, 8.6, 5.4)
card(s, 9.3, 1.5, 3.55, 2.6, "1 mm (reliable run)", ["FEXT and NEXT within a few dB.", "Best-fit δ = 0.23 % (rms 1.9 dB): supports the signal-to-own-grounds value."])
card(s, 9.3, 4.3, 3.55, 2.6, "5 mm", ["FEXT envelope agrees.", "NEXT > 25 GHz: 3-D has no nulls, 6–10 dB higher; it sits on the run's error floor (−56 dB)."], hc=RED)

# 9 length
s = new(); title(s, "Crosstalk vs electrode length", "Length enters only through e^{−γL}: one model for every length")
figp(s, os.path.join(RI, "length.png"), 0.5, 1.5, 12.35, 3.9)
card(s, 0.5, 5.6, 4.0, 1.3, "FEXT", ["grows ∝ f·L·Δn, then flattens with loss"])
card(s, 4.67, 5.6, 4.0, 1.3, "NEXT", ["oscillates with L; peak level ≈ δ/2"])
card(s, 8.84, 5.6, 4.01, 1.3, "EO crosstalk", ["peaks where the aggressor's own walk-off null falls (60 GHz, ~14.7 mm)"], hc=RED)

# 10 GUI
s = new(); title(s, "Crosstalk window: the four graphs", "14 mm, n_{g} = 2.27, source 50 Ω, load 40 Ω; band = δ from 0.21 to 0.61 %")
figp(s, os.path.join(RI, "gui4.png"), 0.5, 1.5, 7.8, 5.4)
card(s, 8.5, 1.5, 4.35, 1.05, "(1) Electrical, all ports 50 Ω", ["|S_{31}| NEXT, |S_{41}| FEXT, as a VNA measures"], size=12)
card(s, 8.5, 2.7, 4.35, 2.05, "(2) Voltage seen by the light", ["Device source/load; A driven, B's input on its idle driver."], size=12)
eq(s, "eq_H", 8.62, 3.75, w=4.1)
card(s, 8.5, 4.9, 4.35, 1.05, "(3) EO crosstalk", [], size=12)
eq(s, "eq_EO", 8.62, 5.3, h=0.55)
card(s, 8.5, 6.1, 4.35, 0.8, "(4) Victim voltages at B's two ends", [], size=12)

# 10b EO derivation 1
s = new(); title(s, "EO response of lines A and B: derivation (1/2)", "Exact; this is what the twin and the INTERCONNECT element compute")
text(s, 0.6, 1.55, 12.2, 0.75, ["**1. Modes decouple.** Both lines see the same terminations (Z_{s} at the input: A's driver, B's idle driver; R_{t} at the far end). The aggressor EMF E on A = E/2 even + E/2 odd, each on its own single line:"], size=14)
eq(s, "eq_Vm", 0.9, 2.35, h=0.5)
eq(s, "eq_Vplus", 0.9, 2.95, h=0.62)
text(s, 0.6, 3.7, 12.2, 0.75, ["**2. What the light integrates.** The optical wave travels at c/n_{g}, so each mode's voltage is averaged in the frame of the light. Closed form, forward (V^{+}) and backward (V^{−}) parts:"], size=14)
eq(s, "eq_Hm", 0.9, 4.45, h=0.62)
text(s, 0.6, 5.25, 12.2, 0.45, ["**3. Back to the lines** (same recombination as for the S-parameters):"], size=14)
eq(s, "eq_HAB", 0.9, 5.75, h=0.62)
text(s, 0.6, 6.55, 12.2, 0.4, "V_{m}^{−} = Γ_{t,m} · exp(−2γ_{m}L) · V_{m}^{+};  β_{o} = 2πf n_{g}/c. Electrical S-parameters (NEXT, FEXT) use the same modal waves with z_{0} = 50 Ω terminations.", size=12, color=MUTED)

# 10c EO derivation 2
s = new(); title(s, "EO crosstalk vs NEXT and FEXT: derivation (2/2)", "First order in δ and Δβ, terminations matched to the line (Z_{s} = R_{t} = Z)")
eq(s, "eq_ratio", 0.6, 1.6, h=0.75)
card(s, 0.5, 2.6, 4.0, 2.75, "Launch term δ/4 (NEXT-like)", ["The impedance split launches unequal even and odd waves: Z_{m}/(Z_{m}+Z_{s}) differs by δ/4.",
     "Result: a forward wave of relative size δ/4 starts on B at the input and travels with the light.", "Same δ/4 as in NEXT ≈ (δ/4)(1 − t^{2})."])
card(s, 4.67, 2.6, 4.0, 2.75, "Accumulated term (FEXT-like)", ["Even and odd waves drift apart at Δβ; the light averages the growing voltage on B.",
     "Velocity-matched, lossless line: average = half the far-end value."])
eq(s, "eq_limit", 4.8, 5.45, h=0.6)
card(s, 8.84, 2.6, 4.01, 2.75, "Backward waves (NEXT path)", ["Travel against the light: their integral is phase-mismatched and suppressed by"], hc=RED)
eq(s, "eq_back", 9.0, 4.1, h=0.65)
text(s, 8.95, 4.85, 3.8, 0.5, "−22 / −31 / −37 dB at 10 / 30 / 60 GHz (14 mm)", size=12, color=MUTED)
table(s, 0.5, 6.15, 8.2, [["14 mm, δ = 0.61 %", "10 GHz", "30 GHz", "60 GHz"],
    ["Exact twin (Z_{s} = R_{t} = Z)", "−53.7 dB", "−46.3 dB", "−36.3 dB"], ["Formula above", "−53.6 dB", "−46.2 dB", "−36.3 dB"]], [3.2, 1.6, 1.6, 1.6], size=12, row_h=0.32)
text(s, 8.95, 6.15, 3.9, 0.95, "Below ~5 GHz, (β + β_{o})L is small and the backward terms add 1–3 dB.", size=12, color=MUTED)

# 11 element
s = new(); title(s, "INTERCONNECT: the \"Coupled TL\" scripted element", "The single-line TL element, computed twice (even, odd) and recombined")
box(s, 0.5, 1.6, 3.6, 2.0, "TL line element (existing)", ["1 table: f, loss, n, Re Z, Im Z", "2 ports + modulation output", "exact V(z) integral, no slicing"])
arrow(s, 4.25, 2.4, 0.8)
box(s, 5.2, 1.6, 3.6, 2.0, "Coupled TL element (new)", ["2 tables: coupled_even/odd.txt", "ports A1, A2, B1, B2", "modulation A, B (+ 2 A, 2 B), far/near end"], fill=RGBColor(0xFD, 0xEC, 0xEE), hc=RED)
card(s, 9.0, 1.6, 3.85, 2.0, "Recombination", ["S own line = ½(S^{e} + S^{o})", "S other line = ½(S^{e} − S^{o})", "same for the modulation outputs"])
card(s, 0.5, 3.85, 6.05, 3.05, "Tables coupled_even.txt / coupled_odd.txt", ["Columns: frequency (Hz), loss (dB/m), microwave index, Re Z_{0}, Im Z_{0} (Ω).",
     "Written by the GUI (\"Export tables\"): 3-D line fit ± the split (Δn × 1.146, δ = 0.61 %).",
     "0.05 GHz to the Nyquist frequency; above 60 GHz, extrapolated with the single-line fit forms."])
card(s, 6.78, 3.85, 6.07, 3.05, "Code check", ["LSF code executed line by line vs Python twin: ≤ 4.5·10^{−14}.",
     "INTERCONNECT run (14 mm): aggressor within 0.09 dB / 0.44°; crosstalk error 72–84 dB below the aggressor."])

# 12 test bench
s = new(); title(s, "INTERCONNECT test bench (PAIR_MOD / FE / NE)", "Isolates the element: driver, idle victim driver, two terminations, analyser")
for i, (h, b) in enumerate([("ENA", ["output → driver", "inputs 1, 2 ← two outputs"]), ("SRC", ["driver, Z_{s}"]), ("CTL", ["Coupled TL element"]), ("LOAD_A2 / B2", ["terminations R_{t}"])]):
    box(s, 0.5 + i * 3.15, 1.65, 2.85, 1.35, h, b)
    if i < 3: arrow(s, 3.38 + i * 3.15, 2.15, 0.25)
card(s, 0.5, 3.25, 6.05, 3.65, "Why this form", ["LOAD_B1 = Z_{s}: victim's input on its own idle driver.",
     "Compound with scattering-data analysis on: only there are reflections exact (Step 0 result).",
     "Three rows: an analyser reads two outputs; MOD = modulation A/B, FE = far end, NE = near end."])
card(s, 6.78, 3.25, 6.07, 3.65, "What is compared", ["The element's outputs are electrical voltages: 'modulation A/B' = H_{A}, H_{B}, the voltages the light integrates (previous slides). The ENA reads them as complex transfer functions.",
     "Scale (2.00: ENA source = half the EMF), FIR latency (341 ps) and phase convention are taken from A; the same corrections are applied to B.",
     "Optical EO response (photodiode): checked for the single line (Step 1, 0.1 dB); for the pair it comes with the full cascade run."], size=12.5)

# 13 full cascade
s = new(); title(s, "Full MZM cascade in crosstalk mode", "Two modulators on one coupled electrode; built and dry-run, first INTERCONNECT run pending")
box(s, 0.5, 1.6, 2.4, 1.1, "Driver A (Z_{s})", ["ENA or PRBS/NRZ"])
box(s, 0.5, 3.0, 2.4, 1.1, "Idle driver B (Z_{s})", [""])
box(s, 3.4, 1.6, 3.2, 2.5, "Coupled TL (CTL_1)", ["A1 ───────── A2", "B1 ───────── B2", "outputs: modulation A, 2 A, B, 2 B"], fill=RGBColor(0xFD, 0xEC, 0xEE), hc=RED)
box(s, 7.1, 1.6, 2.2, 1.1, "Load A (R_{t})", [""]); box(s, 7.1, 3.0, 2.2, 1.1, "Load B (R_{t})", [""])
arrow(s, 2.95, 1.95, 0.4); arrow(s, 2.95, 3.35, 0.4); arrow(s, 6.65, 1.95, 0.4); arrow(s, 6.65, 3.35, 0.4)
box(s, 9.6, 1.6, 3.25, 1.1, "MZM A → PIN_A", ["ENA input 1 / EYE_1"], hc=GREEN)
box(s, 9.6, 3.0, 3.25, 1.1, "MZM B → PIN_B", ["ENA input 2 / EYE_2"], hc=GREEN)
card(s, 0.5, 4.45, 6.05, 2.45, "What it gives", ["ENA: B/A = EO crosstalk in the receiver, checked against the twin.",
     "Eyes: A's eye on the coupled line; EYE_2 = crosstalk reaching B (driver idle).",
     "Optics, bias, drive, source/load from the sidebar."])
card(s, 6.78, 4.45, 6.07, 2.45, "Reading the ENA", ["The raw ENA trace is a flat level with a fast ripple: a delay-free spike of the optical path, ~10× the response.",
     "The tool removes it (time gate + cos(πf/f_{s}) correction) and saves the cleaned EO response and B/A."], hc=RED)

# 14 limits & next
s = new(); title(s, "Limits and next steps")
card(s, 0.5, 1.4, 6.05, 5.5, "Limits", ["**δ band** 0.21–0.61 %: NEXT uncertain by ~9 dB; FEXT and EO only ±1 dB.",
     "**Port / pad coupling** not modelled (Istanbul 3-D: port FEXT term ≈ −43 dB, dominant on short lines).",
     "**5 mm NEXT** above 25 GHz unresolved (3-D error floor).", "**Straight electrode** only in the cascade; data to 60 GHz.",
     "**200 µm pitch** uniform line only."], hc=RED)
card(s, 6.78, 1.4, 6.07, 5.5, "Next steps", ["**1.** First INTERCONNECT run of the full cascade (ENA, then eyes).",
     "**2.** Narrow δ: second reliable 3-D length (long line, fine symmetric mesh) → keep 0.21–0.25 % if confirmed.",
     "**3.** Lab: check that probes/pads tie each line's grounds.", "**4.** Victim driven with its own pattern: combined eye penalty.",
     "**5.** Slow-wave (Istanbul) line: split from 3-D lengths or a matched eigenmode cell."])

prs.save(os.path.join(HERE, "..", "crosstalk_meeting.pptx")); print("saved", len(prs.slides._sldIdLst), "slides")
