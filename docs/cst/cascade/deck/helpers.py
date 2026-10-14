"""CST straight line: port artefact, de-embedding, line extraction (company template)."""
import os, re, pickle
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
HERE = os.path.dirname(os.path.abspath(__file__)); IMG = os.path.join(HERE, "img")
NAVY = RGBColor(0x0E, 0x22, 0x68); RED = RGBColor(0xF2, 0x00, 0x13); INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x60, 0x6E); CARD = RGBColor(0xF1, 0xF3, 0xF8); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57); BLUE = RGBColor(0x2F, 0x6F, 0xD0); ROW = RGBColor(0xF6, 0xF7, 0xFA)
FONT = "Arial"

prs = Presentation(os.path.join(HERE, "..", "deck2", "build", "template.pptx"))
lst = prs.slides._sldIdLst
for s in list(lst):
    prs.part.drop_rel(s.get(qn("r:id"))); lst.remove(s)
L_TITLE, L_BLANK = prs.slide_layouts[11], prs.slide_layouts[12]


def runs(p, t, size, color, bold=False, italic=False):
    for tok in re.split(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})", t):
        if not tok:
            continue
        b, base, s = bold, 0, tok
        if tok.startswith("**"): b, s = True, tok[2:-2]
        elif tok.startswith("_{"): base, s = -25000, tok[2:-1]
        elif tok.startswith("^{"): base, s = 30000, tok[2:-1]
        r = p.add_run(); r.text = s
        f = r.font; f.name, f.size, f.bold, f.italic = FONT, Pt(size), b, italic; f.color.rgb = color
        if base: r._r.get_or_add_rPr().set("baseline", str(base))


def text(s, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, sa=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(sa)
        sz, col, bd = size, color, bold
        if isinstance(para, tuple):
            para, o = para; sz, col, bd = o.get("size", sz), o.get("color", col), o.get("bold", bd)
        runs(p, para, sz, col, bd, italic)
    return tb


def title(s, t, sub=None):
    text(s, 0.55, 0.42, 10.6, 0.7, t, size=28)
    if sub:
        text(s, 0.6, 1.02, 11.8, 0.4, sub, size=15, color=NAVY, bold=True)


def rect(s, x, y, w, h, fill=CARD, line=None, rounded=True):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded: sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def pic(s, name, x, y, w=None, h=None):
    path = os.path.join(IMG, name); W, H = Image.open(path).size
    if w and h:                         # fit inside the box, centred
        sc = min(w / W, h / H); ww, hh = W * sc, H * sc
        return s.shapes.add_picture(path, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def label(s, x, y, w, t, color=MUTED):
    text(s, x, y, w, 0.3, t, size=11, color=color, bold=True)


def table(s, x, y, w, rows, col_w, size=11, row_h=0.36, left=(0,), hl_rows=()):
    gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))); tb = gt.table
    pr = tb._tbl.tblPr
    for k in ("bandRow", "firstRow"): pr.set(k, "0")
    st = pr.find(qn("a:tableStyleId"))
    if st is not None: pr.remove(st)
    tot = sum(col_w)
    for j, cw in enumerate(col_w): tb.columns[j].width = Inches(w * cw / tot)
    for i, r in enumerate(rows):
        tb.rows[i].height = Inches(row_h)
        for j, v in enumerate(r):
            c = tb.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE; p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.LEFT if j in left else PP_ALIGN.CENTER
            head = i == 0
            col = WHITE if head else (RED if (v.startswith("●") ) else INK)
            runs(p, str(v), size, col, bold=head or j == 0)
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if head else (RGBColor(0xFD, 0xEC, 0xEE) if i in hl_rows else (ROW if i % 2 else WHITE))
    return gt



def new():
    return prs.slides.add_slide(L_BLANK)
