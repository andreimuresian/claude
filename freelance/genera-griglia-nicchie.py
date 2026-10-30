# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/home/user/claude/freelance/griglia-nicchie.xlsx"

ARIAL   = "Arial"
BLUE    = Font(name=ARIAL, size=11, color="0000FF")           # celle da compilare
BLACK   = Font(name=ARIAL, size=11)
GREY    = Font(name=ARIAL, size=11, color="7C8A91")
BOLD    = Font(name=ARIAL, size=11, bold=True)
H1      = Font(name=ARIAL, size=15, bold=True, color="141C22")
H2      = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
NOTE    = Font(name=ARIAL, size=10, color="4A565E", italic=True)
LINK    = Font(name=ARIAL, size=10, color="0563C1", underline="single")
HEADFIL = PatternFill("solid", fgColor="0E6E5E")
WARNFIL = PatternFill("solid", fgColor="FFF2CC")
THIN    = Border(bottom=Side(style="thin", color="D3D0C7"))
WRAP    = Alignment(wrap_text=True, vertical="top")

# criterio, descrizione, peso, come si valuta (1 = peggio, 5 = meglio)
CRITERI = [
    ("Bisogno urgente / obbligo", "C'è una scadenza di legge o un costo evidente che spinge ad agire ora?",
     20, "1 = nice-to-have; 5 = obbligo di legge o costo alto e documentato"),
    ("Bisogno permanente", "Il bisogno esiste anche se non cambia nessuna legge e non c'è nessuna scadenza?",
     20, "1 = nasce da una scadenza o da un'ondata normativa; 5 = problema operativo o di cassa che esiste sempre"),
    ("Budget del cliente", "I clienti tipo hanno ricavi >1 M€ e capacità di spesa?",
     20, "1 = micro-imprese; 5 = PMI con >1 M€ di ricavi e abitudine a pagare software"),
    ("Ricavi ricorrenti", "Il cliente paga ogni anno o ogni mese?",
     10, "1 = una tantum; 5 = abbonamento naturale"),
    ("Scalabilità", "Quanto costa servire un cliente in più? Margine?",
     15, "1 = molto lavoro per cliente; 5 = costo marginale quasi zero"),
    ("Canale di vendita", "Sai dove trovare i clienti e come raggiungerli?",
     10, "1 = nessun canale chiaro; 5 = canale noto (associazioni, rivenditori, consulenti)"),
    ("Concorrenza (5 = poca)", "Quanti concorrenti e quanto forti?",
     10, "1 = mercato saturo; 5 = pochi o nessuno ben posizionato"),
    ("Vantaggio tuo", "Capisci il problema meglio di chi lo risolve oggi?",
     10, "1 = nessun vantaggio; 5 = lo conosci dall'interno"),
    ("Testabile a poco prezzo", "Puoi verificare la domanda con un prototipo e conversazioni, senza spendere?",
     5, "1 = serve capitale; 5 = test in 2 settimane a costo zero"),
]

TIPI_SEGNALE = ["Obbligo di legge", "Incentivi / fondi", "Annunci di lavoro", "Recensioni negative",
                "Acquisizioni di verticali", "Forum / community", "Ritardi di pagamento",
                "Carenza di manodopera", "Dolore operativo permanente", "Altro"]
STATI = ["Da studiare", "In studio", "Da intervistare", "Intervistata", "Scartata", "Scelta"]

# segmenti B2B: nome, quota stimata (dall'altro chatbot), n. indicativo
SEGMENTI = [
    ("Manifattura avanzata e meccanica", "26–28%", "11.000–11.500"),
    ("Commercio all'ingrosso B2B", "20–22%", "8.500–9.000"),
    ("Servizi corporate, ICT e software", "12–14%", "5.000–5.800"),
    ("Costruzioni, impiantistica, real estate B2B", "10–12%", "4.200–4.800"),
    ("Chimica, polimeri, gomma-plastica", "7–8%", "3.000–3.300"),
    ("Logistica, spedizioni e trasporti B2B", "6–7%", "2.500–2.900"),
    ("Farmaceutica, biotech, dispositivi medici", "4–5%", "1.600–2.000"),
    ("Energia, ambiente, utilities B2B", "3–4%", "1.300–1.600"),
    ("Telecomunicazioni e reti", "1,5–2%", "600–800"),
]

# lista lunga: nicchia, cliente tipo, tipo di segnale, fonte del segnale, note (stato di verifica, date)
V  = "Segnale verificato in ricerca (stampa di settore)."
IP = "Solo ipotesi nostra: domanda non verificata."
DV = "Segnale da verificare: non cercato."
CANDIDATE = [
    # ---- obblighi di legge con scadenza
    ("NIS2: preparazione alla conformità per PMI fornitrici di soggetti critici",
     "PMI in filiere critiche; consulenti e MSP come canale", "Obbligo di legge", "Noze; Agenda Digitale",
     V + " Notifica incidenti operativa dal 15/01/2026; misure di base con scadenze da ottobre 2026 per chi ha ricevuto la comunicazione ad aprile 2025. Mercato affollato di consulenti. Numero di imprese coinvolte non trovato."),
    ("NIS2 per MSP e consulenti: piattaforma multi-cliente per gestire la conformità di più aziende",
     "MSP, studi di consulenza IT", "Obbligo di legge", "Noze; Agenda Digitale",
     V + " Variante della precedente: vendi al canale, non alla PMI finale. Da confrontare con piattaforme esistenti."),
    ("Cyber Resilience Act: gestione segnalazione vulnerabilità e fascicolo tecnico per produttori",
     "Produttori di prodotti con elementi digitali (elettronica, IoT, macchine con software)", "Obbligo di legge",
     "Agenda Digitale; IPSOA",
     V + " Obbligo di segnalazione dall'11/09/2026 (preallarme in 24 ore); piena applicazione 11/12/2027; sanzioni fino a 15 M€ (stampa). Vantaggio tuo: medio-alto."),
    ("Direttiva RED (cybersicurezza): assistente alla conformità EN 18031 per produttori di apparecchi radio/IoT",
     "Produttori e importatori di apparecchi radio e dispositivi connessi", "Obbligo di legge", "TÜV SÜD; Yarix",
     V + " Applicabile dal 01/08/2025. Stato di armonizzazione degli standard discordante tra le fonti. Vantaggio tuo: alto (RF)."),
    ("Regolamento Macchine 2023/1230: generatore di fascicolo tecnico e dichiarazione UE",
     "Costruttori di macchine (Lombardia è un polo)", "Obbligo di legge", "Certifico; UCIMU",
     V + " Applicazione dal 20/01/2027; UCIMU chiede una proroga condizionata. Dettagli sulla documentazione tecnica non verificati."),
    ("AI Act: inventario dei sistemi di IA e conformità agli obblighi di trasparenza (art. 50)",
     "PMI che usano o forniscono sistemi di IA", "Obbligo di legge", "Agenda Digitale; Noze",
     V + " Trasparenza dal 02/08/2026; alto rischio rinviato a dicembre 2027 / agosto 2028 (accordo provvisorio del 07/05/2026, da verificare in Gazzetta UE). Domanda specifica debole."),
    ("Trasparenza retributiva: annunci e report sul divario di genere",
     "Aziende con 100+ dipendenti, uffici HR", "Obbligo di legge", "QuiFinanza; PartitaIVA.it",
     V + " Stato del decreto di recepimento discordante tra le fonti (D.Lgs. 96/2026?). Report annuale dal 07/06/2027 per 250+ dipendenti. Il target è piccolo: poche aziende con 100+ dipendenti."),
    ("PPWR imballaggi: gestione dati e dichiarazioni di conformità",
     "Produttori e utilizzatori di imballaggi (Lombardia ha molti)", "Obbligo di legge", "Assolombarda",
     V + " Applicabile dal 12/08/2026; date di etichettatura discordanti (2026, 2027, 2028). Assolombarda ha attivato un desk dedicato."),
    ("Passaporto digitale di prodotto: raccolta dati per batterie e tessile",
     "Produttori di batterie e tessile", "Obbligo di legge", "Bluestone PIM; PMI.it",
     V + " Batterie dal 18/02/2027; tessile non ancora definito. Dominato da grandi fornitori e dagli atti delegati."),
    ("Accessibilità digitale (EAA): audit e correzione per e-commerce",
     "E-commerce italiani", "Obbligo di legge", "Agenda Digitale; Iubenda",
     V + " Obbligo dal 28/06/2025; piattaforma di segnalazione AgID dal 11/03/2026; sanzioni da verificare (nessun procedimento pubblico trovato a giugno 2026). Molti fornitori."),
    # ---- incentivi / fondi
    ("Iperammortamento 2026: pratica per perizia asseverata, certificazione e comunicazioni al GSE",
     "Studi di ingegneria che firmano perizie; imprese manifatturiere", "Incentivi / fondi",
     "Agenda Digitale; Pasquale Pillitteri",
     V + " Piattaforma GSE dal 12/06/2026; comunicazione di completamento entro 15/11/2028. La perizia è di un ingegnere iscritto all'albo: collegato agli studi che stai contattando. Regole (soglie, SaaS) discordanti tra le fonti."),
    ("Fibra ottica: report di collaudo e documentazione per imprese di posa",
     "Imprese di posa fibra e installatori", "Incentivi / fondi", "Domani",
     V + " Piano Italia 1 Giga in ritardo: 250–700 mila civici potrebbero uscire dal Piano. Committenti con portali propri; strumenti degli OTDR con report già inclusi."),
    ("Comunità energetiche: piattaforma di gestione per CER, condomini e comuni",
     "Promotori di CER, amministratori di condominio, comuni", "Incentivi / fondi", "InfoBuild Energia",
     V + " Obiettivo di 15.000 CER a giugno 2026 probabilmente mancato; regole ancora instabili."),
    ("Diagnosi energetica obbligatoria: strumento per redigere e gestire le diagnosi",
     "Studi di ingegneria, ESCO, grandi imprese ed energivori", "Incentivi / fondi", "da cercare",
     DV + " Obbligo ricorrente (ogni 4 anni per le imprese soggette). Da verificare sulle norme e sul mercato."),
    # ---- telecom / reti
    ("FWA: pianificazione e monitoraggio per piccoli operatori wireless",
     "Piccoli provider FWA/WISP", "Altro", "Segugio (dati AGCOM)",
     V + " Oltre 2,7 milioni di linee FWA, +13% in un anno; frequenze 26 GHz in assegnazione. Nessun dato recente sul numero di piccoli provider."),
    ("Audit e ottimizzazione dei costi telecom per PMI",
     "PMI con molte linee e contratti", "Altro", "da cercare",
     IP + " Vantaggio tuo alto; modello a risparmio condiviso. Da verificare quanto si risparmia davvero."),
    # ---- verticali acquisite dai grandi gruppi
    ("Pianificazione turni e presenze per magazzini, logistica e conto terzi",
     "Logistica, confezionamento conto terzi, magazzini", "Acquisizioni di verticali", "Aventis Advisors (Zucchetti)",
     "Zucchetti ha acquisito la maggioranza di AWMS (personale per fabbriche e logistica): segnale che il verticale vale. In zona: 6 imprese di conto terzi nell'elenco. Software di turni già presenti."),
    ("Gestione commesse e avanzamento lavori per piccole imprese edili e impiantisti",
     "Imprese edili, impiantisti", "Acquisizioni di verticali", "Aventis Advisors (TeamSystem)",
     "TeamSystem ha fatto 7 acquisizioni nel comparto industria/costruzioni/real estate: segnale di valore del verticale. Mercato presidiato da gruppi grandi."),
    ("Scadenzario obblighi di sicurezza sul lavoro (formazione, visite, DPI) per PMI",
     "PMI manifatturiere e di servizi", "Acquisizioni di verticali", "Aventis Advisors (TeamSystem: legal, compliance, sicurezza)",
     "5 acquisizioni TeamSystem in legal/compliance/sicurezza. Prodotti già esistenti; da mappare i concorrenti."),
    # ---- idee nostre, domanda da verificare
    ("Report di collaudo automatici da dati di misura per laboratori e officine",
     "Laboratori di prova, officine di precisione, assemblatori", "Altro", "idea precedente",
     IP + " Gli strumenti hanno già software di report. Vantaggio tuo: alto."),
    ("Scadenzario tarature degli strumenti di misura (ISO 9001)",
     "Officine, laboratori, PMI certificate", "Altro", "idea precedente", IP),
    ("Preventivatore con codici articolo per distributori tecnici",
     "Distributori di bulloneria, acciai, plastiche, macchine utensili", "Altro", "idea precedente",
     IP + " Probabile presenza di gestionali; da verificare."),
    ("Estrazione dati da ordini, DDT e preventivi in PDF per il back-office",
     "PMI con ufficio ordini", "Annunci di lavoro", "da cercare",
     DV + " Da cercare tra gli annunci per ruoli ripetitivi (ufficio ordini, data entry)."),
    ("Assistente a documentazione tecnica per l'assistenza post-vendita dei costruttori di macchine",
     "Costruttori di macchine", "Altro", "idea Gemini", IP),
    ("Automazione di relazioni, perizie e pratiche per studi tecnici",
     "Studi di ingegneria e professionisti", "Altro", "studi contattati",
     IP + " Gli studi (DEAL, Parodi, Lommano) fanno pratiche antincendio, certificazioni energetiche, perizie."),
    ("Gestione dei file di misura RF e ottici per R&D e produzione",
     "Laboratori R&D e produzione elettronica", "Altro", "tuo campo", IP + " Vantaggio tuo: alto."),
    ("Schede di sicurezza e conformità CLP/REACH per chimica e gomma-plastica",
     "Aziende chimiche e della gomma-plastica", "Altro", "segmenti B2B (non verificati)", DV),
    ("Documentazione elettronica di lotto (batch record) per piccole aziende farmaceutiche/CDMO",
     "Piccole aziende farmaceutiche e CDMO", "Altro", "segmenti B2B (non verificati)",
     DV + " Barriera regolatoria alta (GMP)."),
    ("Formazione e simulazioni di phishing per PMI (obblighi di formazione NIS2)",
     "PMI", "Obbligo di legge", "Noze; Agenda Digitale", V + " Mercato affollato e a basso prezzo."),
    ("Manutenzione predittiva a basso costo per macchine di PMI manifatturiere",
     "PMI manifatturiere", "Altro", "idea nostra", IP + " Concorrenza forte (grandi fornitori)."),
    ("Incassi e solleciti: gestione automatica del credito e previsione di cassa per PMI B2B",
     "PMI B2B (ingrosso, manifattura) con molte fatture", "Ritardi di pagamento", "Cribis; Cerved; Intrum",
     V + " Cribis giugno 2026: solo il 42% delle imprese paga puntuale (piccole 41,8%, medie 35,7%, grandi 21,1%). Intrum 2024: ~81 giorni di incasso medi in Italia contro ~50 in Europa. Bisogno di cassa permanente. Concorrenza alta (gestionali, fintech)."),
    ("TMS leggero per piccole imprese di autotrasporto (DDT/CMR, scadenze, pianificazione)",
     "Autotrasportatori conto terzi", "Dolore operativo permanente", "Uomini e Trasporti; Istat",
     V + " 67.349 imprese attive nel 2025 (-22% dal 2015), le ditte individuali -41%: il settore si concentra, i clienti diminuiscono. Il cugino della tua ragazza ha una ditta di trasporti: possibile primo contatto."),
    ("Spedizionieri e magazzini: documenti, tracciamento e dogane",
     "Spedizionieri, operatori logistici", "Dolore operativo permanente", "Istat",
     "Mercato medio (ATECO 52.29 e 52.10). " + DV),
    ("Gestione commesse e tracciabilità per officine e subfornitura meccanica",
     "Officine meccaniche e carpenterie con 10+ addetti", "Dolore operativo permanente", "Istat",
     IP + " Richiede conoscenza del mestiere; presenti ERP e MES."),
    ("Export: documentazione per PMI esportatrici (origine, dogana)",
     "PMI esportatrici", "Dolore operativo permanente", "Istat-ICE 2025",
     DV + " 120.170 imprese esportatrici attive nel 2023; nessun dato su costo della documentazione."),
    ("Gestione documentale per sicurezza prodotti (GPSR) di importatori e venditori online",
     "Importatori e e-commerce di prodotti per consumatori", "Obbligo di legge", "da cercare", DV),
]

FONTI_URL = {
    "NIS2": ["https://www.noze.it/insights/nis2-pmi-cyberagent/",
             "https://www.agendadigitale.eu/sicurezza/nis2-scadenze-2026-la-guida-per-le-piccole-imprese/"],
    "AI Act": ["https://www.agendadigitale.eu/industry-4-0/digital-omnibus-e-ai-act-cosa-puo-slittare-e-cosa-va-fatto-prima-di-agosto-2026/",
               "https://www.noze.it/insights/ai-act-digital-omnibus-2026/"],
    "Cyber Resilience Act": ["https://www.agendadigitale.eu/sicurezza/cyber-resilience-act-la-scadenza-chiave-arriva-l11-settembre-2026/",
                             "https://www.ipsoa.it/documents/quotidiano/2026/09/02/cyber-resilience-act-11-settembre-2026-scatta-obbligo-segnalare-vulnerabilita-incidenti-gravi"],
    "Direttiva RED": ["https://www.tuvsud.com/en-us/services/product-certification/red-cybersecurity-requirements",
                      "https://www.yarix.com/Advisory-and-Consulting-Services/RED-Compliance"],
    "Regolamento Macchine": ["https://certifico.com/macchine/regolamento-ue-2023-1230",
                             "https://www.ilsussidiario.net/news/regolamento-macchine-ue-rosa-ucimu-mancano-criteri-comuni-e-chiari-serve-una-proroga/2986256/amp/"],
    "Trasparenza retributiva": ["https://quifinanza.it/lavoro/trasparenza-salariale-7-giugno-2026/960990/",
                                "https://www.partitaiva.it/trasparenza-salariale/"],
    "PPWR": ["https://www.assolombarda.it/servizi/ambiente/informazioni/ppwr-dal-12-agosto-in-applicazione-il-nuovo-regolamento-ue-sugli-imballaggi"],
    "Passaporto digitale di prodotto": ["https://www.bluestonepim.com/blog/digital-product-passport-for-batteries",
                                        "https://www.pmi.it/tecnologia/prodotti-e-servizi-ict/492107/passaporto-digitale-prodotto.html"],
    "Accessibilità (EAA)": ["https://www.agendadigitale.eu/mercati-digitali/ecommerce/siti-ecommerce-accessibili-guida-in-5-step-per-evitare-sanzioni/",
                            "https://www.iubenda.com/it/help/178476-accessibilita-101/"],
    "Iperammortamento 2026": ["https://www.agendadigitale.eu/industry-4-0/iperammortamento-2026-la-guida-cose-e-come-fare-le-comunicazioni-al-gse/",
                              "https://pasqualepillitteri.it/perizia-iperammortamento"],
    "Acquisizioni TeamSystem / Zucchetti": ["https://aventis-advisors.com/teamsystem-acquisitions/"],
    "Piano Italia 1 Giga": ["https://www.editorialedomani.it/economia/banda-ultralarga-cantieri-obiettivi-pnrr-ritardi-regioni-oqbs3i02",
                            "https://tariffe.segugio.it/news-tariffe/italia-1-giga-nuovi-ritardi-per-la-banda-ultra-larga-cosa-sta-succedendo-00043657.aspx"],
    "FWA e frequenze 26 GHz": ["https://tariffe.segugio.it/news-tariffe/internet-senza-fibra-la-fwa-supera-2-7-milioni-di-linee-00047779.aspx",
                               "https://www.smartworld.it/news/frequenze-5g-26-ghz-gara-italia-agosto-2026.html"],
    "Fase 2 — numeri di mercato": [
        "https://esploradati.istat.it (dataset: Classe di addetti, settori economici (Ateco 5 cifre) - Italia, anno 2024)",
        "https://accredia.it/wp-content/uploads/2024/07/Cartella-stampa-breve-dati-al-31-12-2025_ITA-1.pdf",
        "https://assomac.it/media/documents/rapporto_federmacchine_2020.pdf",
        "https://anie.it/federazione/chi-siamo/rapporto-annuale-2025_def",
        "https://www.cni.it/images/eventi/2021/01_Lelaborazione_delle_perizie_13lug2021.pdf",
        "https://www.fibertelecom.com/wt/2022/PDF/02-puglia/FTWWT2022-puglia-03.pdf"],
    "Comunità energetiche (CER)": ["https://www.infobuildenergia.it/approfondimenti/cer-comunita-energetiche-in-italia/"],
}

# ---- Fase 2: stime di mercato (Italia). nicchia (inizio del nome in Lista lunga), etichetta breve,
# aziende bassa/base/alta, fonte, affidabilità, prezzo annuo bassa/base/alta (ipotesi), lettura
FASE2 = [
    ("NIS2: preparazione", "NIS2 per PMI", 7000, 13500, 30000,
     "Istat ASIA 2024: 13.574 unità locali con 50+ addetti nei settori NIS2 (perimetro mio, approssimato). Stampa: 20.000-50.000 enti. Sono unità locali: limite superiore", "Media",
     600, 1500, 4000, "Base ampia, ma mercato affollato di consulenti: serve un canale forte."),
    ("NIS2 per MSP", "NIS2 per MSP e consulenti", 2700, 4400, 8000,
     "Istat 2024: ATECO 62.02/62.03/62.09 e 63.11 con 10+ addetti (unità locali); alto = tutto 62 e 63.1 con 10+", "Media",
     1200, 2400, 6000, "Vendi al canale: meno clienti ma con più aziende ciascuno. Da confrontare con le piattaforme esistenti."),
    ("Cyber Resilience Act", "Cyber Resilience Act", 1700, 4400, 17000,
     "Istat 2024: ATECO 26 (1.692), 26+27 (4.396), + 28 e software 62.01 con 10+ (17.011); solo una parte produce prodotti con elementi digitali", "Media-bassa",
     1500, 3000, 6000, "Norma europea: il mercato vero è più grande dell'Italia, ma va servito in inglese e con concorrenti internazionali."),
    ("Direttiva RED", "Direttiva RED (radio/IoT)", 250, 780, 1700,
     "Istat 2024: ATECO 26.3 (781 unità locali, 254 con 10+); alto = ATECO 26 con 10+ (1.692). Gli importatori non sono inclusi", "Media-bassa",
     2000, 4000, 8000, "Troppo piccolo come solo mercato italiano; ha senso solo come prodotto europeo."),
    ("Regolamento Macchine", "Regolamento Macchine", 2200, 9000, 22700,
     "Istat 2024: ATECO 28 = 22.720 unità locali, 9.059 con 10+, 2.179 con 50+. Controllo: Federmacchine 2020, 5.050 imprese associate", "Media",
     1500, 2500, 5000, "Mercato ben definito e vicino a te; la scadenza (20/01/2027) può slittare."),
    ("Iperammortamento 2026", "Iperammortamento 2026 (perizie)", 5000, 12000, 20000,
     "CNI 2020: 4,5%-9% degli ingegneri ha firmato perizie, in media 10 ciascuno (+ ordine di grandezza mio)", "Bassa",
     300, 600, 1500, "Tanti utenti ma prezzo basso: serve come aggiunta a un prodotto più grande."),
    ("Fibra ottica", "Fibra ottica (posa)", 400, 800, 6600,
     "Istat 2024: ATECO 42.22 = 805 unità locali (401 con 10+); alto = + 43.21 con 10+ (5.781). La posa è distribuita su più codici", "Media-bassa",
     1000, 2000, 4000, "Mercato piccolo e dipendente da pochi committenti."),
    ("FWA: pianificazione", "FWA (piccoli operatori)", 230, 600, 1300,
     "Istat 2024: ATECO 61.10 = 1.315 e 61.20 = 226 unità locali. I WISP non sono isolabili; base = mio giudizio", "Bassa",
     800, 1500, 3000, "Troppo piccolo come mercato solo italiano."),
    ("Report di collaudo", "Report di collaudo", 1000, 2000, 3900,
     "Istat 2024: ATECO 71.20 = 9.362 unità locali, 1.041 con 10+; 71.20.1 prodotti = 3.867. Accredia: 1.448 laboratori di prova", "Media",
     600, 1200, 3000, "Mercato di nicchia; gli strumenti di misura hanno già report inclusi."),
    ("Incassi e solleciti", "Incassi e solleciti (cassa)", 5800, 20000, 45700,
     "Istat 2024: unità locali con 10+ addetti in ingrosso (46: 20.256), metallo (25: 16.356), macchinari (28: 9.059); basso = 50+ (5.804). Cribis 2026: solo 42% paga puntuale", "Media",
     600, 1200, 3000, "Bisogno permanente e documentato; concorrenza alta (gestionali e fintech): serve un canale."),
    ("TMS leggero", "TMS per autotrasporto", 900, 9000, 22700,
     "Istat 2024: ATECO 49.41 = 63.633 unità locali, 9.037 con 10+, 915 con 50+; alto = 10+ più 25% delle micro (mia ipotesi). Imprese attive in calo del 22% dal 2015", "Media-bassa",
     500, 1200, 3000, "Bisogno permanente, ma il settore si concentra e i clienti diminuiscono."),
    ("Spedizionieri e magazzini", "Spedizionieri e magazzini", 850, 4100, 6600,
     "Istat 2024: ATECO 52.29 e 52.10 con 10+ = 4.123; 50+ = 849; alto = 52.2 con 10+ (6.583)", "Media",
     1000, 2400, 6000, "Mercato medio con documenti ripetitivi; concorrenza di software di settore."),
    ("Gestione commesse", "Commesse e tracciabilità (officine)", 5000, 16400, 25400,
     "Istat 2024: ATECO 25 con 10+ = 16.356 (25.62 meccanica generale: 4.975); alto = + 28 con 10+", "Media",
     1000, 2400, 5000, "Quota necessaria bassa, ma serve conoscere il mestiere e ci sono già ERP e MES."),
    ("Preventivatore", "Preventivatore per distributori tecnici", 2700, 8000, 13000,
     "Istat 2024: ingrosso di macchinari (466) e prodotti specializzati (467) con 10+ = 8.028; alto = + 25.62 (4.975)", "Media",
     800, 1800, 4000, "Legato ai ricavi del cliente; probabili gestionali già presenti, quindi come aggiunta."),
    ("Assistente a documentazione", "Assistenza post-vendita (macchine)", 2200, 9000, 22700,
     "Istat 2024: ATECO 28: 9.059 con 10+, 2.179 con 50+, 22.720 totali", "Media",
     1500, 3000, 6000, "Stesso mercato delle macchine, ma bisogno permanente: ogni macchina venduta richiede manuali, ricambi e assistenza per anni."),
    ("Pianificazione turni", "Turni per logistica e magazzini", 850, 4100, 7200,
     "Istat 2024: magazzinaggio e supporto ai trasporti (52) con 10+ = 7.223; 52.10 e 52.29 con 10+ = 4.123", "Media",
     1500, 3000, 6000, "Software di turni già diffusi; conta quanto sono complesse le regole del cliente."),
    ("Estrazione dati", "Estrazione dati da ordini e DDT", 5800, 20000, 45700,
     "Istat 2024: come per gli incassi (ingrosso, metallo e macchinari con 10+)", "Media",
     800, 1800, 4000, "Utile in tutti i settori ma generico: va scelto un solo tipo di documento."),
]

DUR = {"NIS2 per PMI": 3, "NIS2 per MSP e consulenti": 3, "Cyber Resilience Act": 3, "Direttiva RED (radio/IoT)": 3,
       "Regolamento Macchine": 3, "Iperammortamento 2026 (perizie)": 1, "Fibra ottica (posa)": 2,
       "FWA (piccoli operatori)": 4, "Report di collaudo": 5, "Incassi e solleciti (cassa)": 5,
       "TMS per autotrasporto": 4, "Spedizionieri e magazzini": 4, "Commesse e tracciabilità (officine)": 5,
       "Preventivatore per distributori tecnici": 5, "Assistenza post-vendita (macchine)": 5,
       "Turni per logistica e magazzini": 5, "Estrazione dati da ordini e DDT": 5}

# punteggi provvisori dopo la Fase 3 (mio giudizio): urgente, permanente, budget, ricorrenti, scalabilità,
# canale, concorrenza (5 = poca), vantaggio, testabile
PUNTEGGI = {
    "Assistente a documentazione": [3, 5, 4, 4, 4, 4, 3, 4, 5],
    "Incassi e solleciti":         [4, 5, 3, 5, 4, 3, 1, 2, 4],
    "Preventivatore":              [3, 5, 3, 4, 2, 3, 3, 3, 3],
    "Regolamento Macchine":        [5, 3, 4, 3, 4, 4, 2, 4, 4],
}


wb = Workbook()

# ============================================================ CRITERI
wc = wb.active
wc.title = "Criteri"
wc["A1"] = "Criteri di valutazione delle nicchie"
wc["A1"].font = H1
wc["A2"] = ("I pesi (in blu) sono una proposta: cambiali e i punteggi della Lista lunga si ricalcolano. "
            "Ogni nicchia riceve da 1 a 5 su ogni criterio.")
wc["A2"].font = NOTE
wc["A2"].alignment = WRAP
wc.merge_cells("A2:D2")
wc.row_dimensions[2].height = 30

heads = ["Criterio", "Domanda", "Peso", "Come si valuta (1 = peggio, 5 = meglio)"]
widths = [30, 60, 8, 70]
for c, (t, w) in enumerate(zip(heads, widths), start=1):
    cell = wc.cell(row=4, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    wc.column_dimensions[get_column_letter(c)].width = w
for i, (nome, dom, peso, come) in enumerate(CRITERI, start=1):
    r = 4 + i
    for c, v in enumerate([nome, dom, peso, come], start=1):
        cell = wc.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
        cell.border = THIN
        cell.font = BLUE if c == 3 else (BOLD if c == 1 else BLACK)
    wc.row_dimensions[r].height = 34
LASTC = 4 + len(CRITERI)
wc.cell(row=LASTC + 1, column=2, value="Somma dei pesi").font = BOLD
wc.cell(row=LASTC + 1, column=3, value="=SUM(C5:C{})".format(LASTC)).font = BOLD

# ============================================================ LISTA LUNGA
wl = wb.create_sheet("Lista lunga")
wl["A1"] = "Lista lunga delle nicchie"
wl["A1"].font = H1
wl["A2"] = ("Una riga per nicchia (lista lunga, Fase 1). Il segnale è solo il punto di partenza: la colonna Note dice "
            "se è stato verificato in ricerca (stampa di settore, non fonti ufficiali) o se è solo un'ipotesi. In blu le celle da "
            "compilare nelle fasi 2-3: stima di aziende target e di prezzo (con fonte) e punteggi 1-5. "
            "'Clienti per 1 M€' e 'Punteggio' si calcolano da soli.")
wl["A2"].font = NOTE
wl["A2"].alignment = WRAP
wl.merge_cells("A2:V2")
wl.row_dimensions[2].height = 34

BASE = ["N", "Nicchia", "Cliente tipo", "Tipo di segnale", "Fonte del segnale", "Aziende target (stima)",
        "Fonte della stima", "Prezzo annuo per cliente (€)", "Clienti per 1 M€ di ricavi", "Concorrenti noti"]
SC0 = len(BASE) + 1                      # prima colonna dei punteggi
HEAD = BASE + [c[0] for c in CRITERI] + ["Punteggio ponderato (1-5)", "Stato", "Note"]
wl_widths = [5, 44, 30, 20, 34, 14, 26, 14, 14, 26] + [11] * len(CRITERI) + [14, 16, 40]

WR = 4   # riga dei pesi
wl.cell(row=WR, column=SC0 - 1, value="Peso →").font = BOLD
for k in range(len(CRITERI)):
    cell = wl.cell(row=WR, column=SC0 + k, value="=Criteri!C{}".format(5 + k))
    cell.font = GREY
HR = 5
for c, (t, w) in enumerate(zip(HEAD, wl_widths), start=1):
    cell = wl.cell(row=HR, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    wl.column_dimensions[get_column_letter(c)].width = w
wl.row_dimensions[HR].height = 48

SA = get_column_letter(SC0)
SZ = get_column_letter(SC0 + len(CRITERI) - 1)
TOT = SC0 + len(CRITERI)
NROWS = 70
dv_tipo = DataValidation(type="list", formula1='"' + ",".join(TIPI_SEGNALE) + '"', allow_blank=True)
dv_stato = DataValidation(type="list", formula1='"' + ",".join(STATI) + '"', allow_blank=True)
dv_pt = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
wl.add_data_validation(dv_tipo)
wl.add_data_validation(dv_stato)
wl.add_data_validation(dv_pt)
for i in range(1, NROWS + 1):
    r = HR + i
    cand = CANDIDATE[i - 1] if i <= len(CANDIDATE) else None
    vals = {1: i}
    if cand:
        vals.update({2: cand[0], 3: cand[1], 4: cand[2], 5: cand[3], TOT + 2: cand[4]})
    for c in range(1, len(HEAD) + 1):
        cell = wl.cell(row=r, column=c, value=vals.get(c))
        cell.border = THIN
        cell.alignment = WRAP
        cell.font = GREY if c == 1 else BLUE
    if cand:
        for key, _lab, lo, ba, hi, src, conf, plo, pba, phi, _l in FASE2:
            if cand[0].startswith(key):
                wl.cell(row=r, column=6, value=ba).font = BLUE
                wl.cell(row=r, column=7, value=src + " (affidabilità: " + conf + ")").font = BLUE
                wl.cell(row=r, column=8, value=pba).font = BLUE
    wl.cell(row=r, column=9, value='=IF(H{0}>0,ROUND(1000000/H{0},0),"")'.format(r)).font = BLACK
    wl.cell(row=r, column=TOT,
            value='=IF(COUNT({a}{r}:{z}{r})=0,"",ROUND(SUMPRODUCT({a}{r}:{z}{r},{a}${w}:{z}${w})/SUM({a}${w}:{z}${w}),2))'
            .format(a=SA, z=SZ, r=r, w=WR)).font = BOLD
    if cand:
        for key, pts in PUNTEGGI.items():
            if cand[0].startswith(key):
                for k, v in enumerate(pts):
                    wl.cell(row=r, column=SC0 + k, value=v).font = BLUE
    wl.cell(row=r, column=TOT + 1, value="Da studiare" if cand else None).font = BLUE
    wl.cell(row=r, column=6).number_format = "#,##0"
    wl.cell(row=r, column=8).number_format = "#,##0"
    wl.cell(row=r, column=9).number_format = "#,##0"
    if cand:
        wl.row_dimensions[r].height = 92
dv_tipo.add("D{}:D{}".format(HR + 1, HR + NROWS))
dv_stato.add("{}{}:{}{}".format(get_column_letter(TOT + 1), HR + 1, get_column_letter(TOT + 1), HR + NROWS))
dv_pt.add("{}{}:{}{}".format(SA, HR + 1, SZ, HR + NROWS))
wl.freeze_panes = wl.cell(row=HR + 1, column=3)
wl.auto_filter.ref = "A{}:{}{}".format(HR, get_column_letter(len(HEAD)), HR + NROWS)

# ============================================================ SEGMENTI B2B
wsg = wb.create_sheet("Segmenti B2B")
wsg["A1"] = "Segmenti B2B lombardi con ricavi sopra 1 M€ — DATI NON VERIFICATI"
wsg["A1"].font = H1
wsg["A2"] = ("Stime fornite da un altro chatbot, senza fonte. Servono come ipotesi di lavoro per scegliere "
             "dove guardare. Da confermare con i dati Istat (unità locali per ATECO e classe di addetti) e della "
             "Camera di Commercio prima di usarle in qualsiasi decisione.")
wsg["A2"].font = NOTE
wsg["A2"].alignment = WRAP
wsg.merge_cells("A2:E2")
wsg.row_dimensions[2].height = 44
for c, (t, w) in enumerate(zip(["Segmento", "Quota stimata", "Aziende (stima)", "Verificato?", "Fonte verificata"],
                               [46, 16, 18, 14, 50]), start=1):
    cell = wsg.cell(row=4, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    wsg.column_dimensions[get_column_letter(c)].width = w
for i, (nome, quota, n) in enumerate(SEGMENTI, start=1):
    r = 4 + i
    for c, v in enumerate([nome, quota, n, "No", None], start=1):
        cell = wsg.cell(row=r, column=c, value=v)
        cell.border = THIN
        cell.font = BLUE if c in (4, 5) else BLACK
        if c in (2, 3, 4):
            cell.fill = WARNFIL

r0 = 4 + len(SEGMENTI) + 2
wsg.cell(row=r0, column=1, value="Altre cifre da verificare").font = BOLD
ALTRE = [
    "Quota di aziende lombarde con ricavi >1 M€ per forma giuridica: ditte individuali <0,5%, società di persone ~17%, società di capitali ~34% (fornite da te, senza fonte).",
    "B2B contro B2C: ~11-12% delle imprese B2B sopra 1 M€ contro ~2% delle B2C (altro chatbot, senza fonte). La direzione è plausibile, il rapporto 5-6x no.",
    "Ripartizione B2B/B2C per forma giuridica: individuali 65/35, persone 55/45, capitali 60/40 (tue stime).",
]
for i, t in enumerate(ALTRE, start=1):
    c = wsg.cell(row=r0 + i, column=1, value=t)
    c.font, c.alignment = BLACK, WRAP
    wsg.merge_cells(start_row=r0 + i, start_column=1, end_row=r0 + i, end_column=5)
    wsg.row_dimensions[r0 + i].height = 32


# ============================================================ FASE 2
w2 = wb.create_sheet("Fase 2", 2)
w2["A1"] = "Fase 2 — dimensione dei mercati e quota necessaria per 1 M€ di ricavi (solo Italia)"
w2["A1"].font = H1
w2["A2"] = ("Numero di aziende: da fonti dove esistono, altrimenti ipotesi mie, con affidabilità indicata. "
            "I PREZZI SONO IPOTESI da validare nelle interviste. Le celle in blu si possono cambiare. "
            "'Quota necessaria' = clienti per 1 M€ diviso aziende target: sopra circa il 10% il mercato italiano è troppo piccolo. 'Bisogno permanente' = 1 se nasce da una scadenza di legge, 5 se è un problema che esiste sempre.")
w2["A2"].font = NOTE
w2["A2"].alignment = WRAP
w2.merge_cells("A2:O2")
w2.row_dimensions[2].height = 44

w2["A4"] = "Margine operativo ipotizzato"
w2["B4"] = 0.70
w2["A5"] = "Tassazione combinata (società + dividendi)"
w2["B5"] = 0.465
w2["A6"] = "Obiettivo netto mensile (€)"
w2["B6"] = 30000
w2["A7"] = "Netto mensile con 1 M€ di ricavi (€)"
w2["B7"] = "=ROUND(1000000*B4*(1-B5)/12,0)"
w2["A8"] = "Ricavi annui necessari per l'obiettivo (€)"
w2["B8"] = "=ROUND(B6*12/(B4*(1-B5)),0)"
for r in range(4, 9):
    w2.cell(row=r, column=1).font = BOLD
    w2.cell(row=r, column=2).font = BLUE if r < 7 else BOLD
w2["B4"].number_format = "0%"
w2["B5"].number_format = "0.0%"
for c in ("B6", "B7", "B8"):
    w2[c].number_format = "#,##0"

h2 = ["Nicchia", "Aziende target (bassa)", "Aziende target (base)", "Aziende target (alta)", "Fonte della stima",
      "Affidabilità", "Prezzo annuo (bassa) €", "Prezzo annuo (base) €", "Prezzo annuo (alta) €",
      "Clienti per 1 M€ (base)", "Quota necessaria (base)", "Quota (caso favorevole)", "Quota (caso sfavorevole)", "Lettura", "Bisogno permanente (1-5, mio giudizio)"]
w2_widths = [34, 12, 12, 12, 52, 12, 12, 12, 12, 12, 12, 12, 12, 60, 14]
HR2 = 10
for c, (t, w) in enumerate(zip(h2, w2_widths), start=1):
    cell = w2.cell(row=HR2, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    w2.column_dimensions[get_column_letter(c)].width = w
w2.row_dimensions[HR2].height = 48
for i, (key, lab, lo, ba, hi, src, conf, plo, pba, phi, lett) in enumerate(FASE2, start=1):
    r = HR2 + i
    vals = [lab, lo, ba, hi, src, conf, plo, pba, phi,
            "=ROUND(1000000/H{0},0)".format(r), "=J{0}/C{0}".format(r),
            "=(1000000/I{0})/D{0}".format(r), "=(1000000/G{0})/B{0}".format(r), lett, DUR.get(lab)]
    for c, v in enumerate(vals, start=1):
        cell = w2.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
        cell.border = THIN
        cell.font = BLUE if c in (2, 3, 4, 7, 8, 9, 15) else (BOLD if c in (1, 11) else BLACK)
    for c in (2, 3, 4, 7, 8, 9, 10):
        w2.cell(row=r, column=c).number_format = "#,##0"
    for c in (11, 12, 13):
        w2.cell(row=r, column=c).number_format = "0.0%"
    w2.row_dimensions[r].height = 58
w2.freeze_panes = w2.cell(row=HR2 + 1, column=2)
r = HR2 + len(FASE2) + 2
w2.cell(row=r, column=1, value="Come leggerla").font = BOLD
for k, t in enumerate([
    "Il numero di aziende viene da Istat 2024 (unità locali, quindi limiti superiori) dove è indicato; le ipotesi residue sono segnate nella colonna Affidabilità. Le norme europee (CRA, RED, Macchine) si applicano in tutta l'UE: il mercato reale può essere molte volte quello italiano, ma richiede un prodotto in inglese e concorrenti internazionali.",
    "Prossimo passo di precisione: il perimetro NIS2 e le quote di prodotti con elementi digitali sono mie approssimazioni, da affinare nelle interviste.",
    "Nessun numero di questo foglio è un dato di mercato verificato per il tuo prodotto: sono ordini di grandezza per scartare le nicchie troppo piccole.",
], start=1):
    c = w2.cell(row=r + k, column=1, value=t)
    c.font, c.alignment = BLACK, WRAP
    w2.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=15)
    w2.row_dimensions[r + k].height = 30

# ============================================================ ISTAT 2024
from openpyxl import load_workbook as _lw
import re as _re
_src = "/home/user/claude/freelance/dati/istat-asia-ateco5-classi-addetti-2024.xlsx"
_wb = _lw(_src, read_only=True, data_only=True)
_rows = {}
for _i, _r in enumerate(_wb.active.iter_rows(values_only=True)):
    if _i < 8 or _r[0] is None:
        continue
    _m = _re.match(r"\s*\[(\w+)\]\s*(.*)", str(_r[0]))
    if not _m:
        continue
    def _n(x):
        try:
            return float(x)
        except (TypeError, ValueError):
            return 0
    _rows[_m.group(1)] = (_m.group(2).strip(), [_n(x) for x in _r[1:11]])
wi = wb.create_sheet("Istat 2024", 3)
wi["A1"] = "Istat ASIA 2024 — unità locali delle imprese attive per classe di addetti (estratto)"
wi["A1"].font = H1
wi["A2"] = ("Fonte: Istat, 'Classe di addetti, settori economici (Ateco 5 cifre) - Italia', anno 2024. "
            "ATTENZIONE: sono UNITÀ LOCALI, non imprese: una società con più sedi conta più volte, quindi sono limiti superiori. "
            "File completo in freelance/dati/.")
wi["A2"].font = NOTE
wi["A2"].alignment = WRAP
wi.merge_cells("A2:L2")
wi.row_dimensions[2].height = 40
hh = ["Codice", "Attività", "UL 0-9", "UL 10-49", "UL 50-249", "UL 250+", "UL totale", "UL con 10+", "UL con 50+",
      "Addetti totali (medi annui)"]
for c, (t, w) in enumerate(zip(hh, [10, 60, 11, 11, 11, 11, 12, 12, 12, 18]), start=1):
    cell = wi.cell(row=4, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    wi.column_dimensions[get_column_letter(c)].width = w
CODICI = ["26", "263", "265", "27", "271", "273", "279", "28", "281", "282", "283", "284", "289",
          "4222", "4321", "6110", "6120", "6190", "6201", "6202", "6203", "6209", "6311", "71201", "71202", "7112"]
for i, cod in enumerate(CODICI, start=1):
    if cod not in _rows:
        continue
    nm, v = _rows[cod]
    r = 4 + i
    vals = [cod, nm, v[0], v[2], v[4], v[6], v[8], v[2] + v[4] + v[6], v[4] + v[6], v[9]]
    for c, x in enumerate(vals, start=1):
        cell = wi.cell(row=r, column=c, value=x)
        cell.font = BLACK
        cell.border = THIN
        if c >= 3:
            cell.number_format = "#,##0"
wi.freeze_panes = "C5"


# ============================================================ FASE 3
COMP = [
    # finalista, concorrente, tipo, cosa offre, prezzo trovato, fonte
    ("Assistenza post-vendita (macchine)", "SYS Suite (Sygest)", "Software italiano",
     "Cataloghi ricambi interattivi 2D/3D, portale web per manuali e offerte, app offline, integrazione ERP/PLM",
     "Non trovato", "https://www.itismagazine.it/featured/cataloghi-3d-e-interattivi-sys-suite-e-la-soluzione"),
    ("Assistenza post-vendita (macchine)", "Certifico Macchine 4 (CEM4)", "Software italiano",
     "Fascicolo tecnico, archivio norme, check list e valutazione dei rischi per il Regolamento 2023/1230",
     "Non trovato", "https://www.certifico.com/certifico-macchine-4/manuale-d-uso-ufficiale-di-certifico-macchine-4-cem4"),
    ("Assistenza post-vendita (macchine)", "Quanos (cataloghi ricambi)", "Software internazionale",
     "Creazione e pubblicazione di cataloghi ricambi", "Non trovato", "https://quanos.com/en/knowhow/spare-parts-catalog/create-and-publish-spare-parts-catalogs/"),
    ("Assistenza post-vendita (macchine)", "Salesforce Field Service", "Software per assistenza sul campo",
     "Gestione interventi e tecnici", "175 €/utente/mese (230 € versione avanzata), contratto annuale",
     "https://www.salesforce.com/it/service/field-service-management/pricing/"),
    ("Assistenza post-vendita (macchine)", "Dynamics 365 Field Service", "Software per assistenza sul campo",
     "Gestione interventi e tecnici", "128,90 $/utente/mese (rivenditore)", "https://gestisoft.com/en/dynamics-365/field-service"),
    ("Assistenza post-vendita (macchine)", "Altri software di field service", "Mercato",
     "Panoramica prezzi", "Mediana circa 168 $/utente/mese (intervallo 39-650 $)", "https://www.costbench.com/software/field-service-management/"),
    ("Assistenza post-vendita (macchine)", "Chatbot (RAG) su manuali: nessuna startup italiana trovata", "Lacuna / rischio",
     "Esistono tesi (IMESA, Padova) e casi internazionali (14.000 tecnici); consulenze su misura possibili",
     "Non applicabile", "https://thesis.unipd.it/handle/20.500.12608/95455"),
    ("Assistenza post-vendita (macchine)", "Servizi di manualistica tecnica", "Servizi",
     "Creazione di manuali e cataloghi ricambi come servizio esterno", "Non trovato", "https://www.reteimprese.it/sers_A105664B100067"),
    ("Incassi e solleciti", "Fatture in Cloud (TeamSystem)", "Gestionale/fatturazione",
     "Solleciti automatici di pagamento, anche per importi sopra una soglia", "Da circa 48-60 € + IVA l'anno (fonti discordanti)",
     "https://www.fattureincloud.it/software-fatturazione/solleciti-pagamento/"),
    ("Incassi e solleciti", "Agicap Italia", "Fintech / tesoreria",
     "Previsione di cassa che raccoglie dati da banche e gestionali; oltre 1.000 clienti in Italia dopo due anni",
     "Non trovato", "https://www.economyup.it/fintech/fintech-e-pmi-come-agicap-italia-automatizza-la-gestione-dei-flussi-di-cassa/"),
    ("Incassi e solleciti", "Previxio", "Fintech / tesoreria",
     "Tesoreria, fatturazione elettronica e previsioni a 12 mesi (da listing non verificato)", "Non trovato",
     "https://www.getapp.com/finance-accounting-software/treasury/f/cash-management/?page=3"),
    ("Incassi e solleciti", "bexio", "Gestionale (Svizzera)", "Solleciti automatici", "Non trovato",
     "https://www.bexio.com/it-CH/software-fatturazione"),
    ("Preventivatore per distributori tecnici", "Salesforce Revenue Cloud (CPQ)", "Enterprise",
     "Configurazione, prezzi e preventivi", "200 €/utente/mese (fascia Advanced, contratto annuale)", "https://www.salesforce.com/it/products/cpq/"),
    ("Preventivatore per distributori tecnici", "ConvertCalculator, iQuoteXpress, Camos CPQ", "CPQ per PMI / prodotti complessi",
     "Preventivi configurabili; Camos per prodotti multi-variante", "Non trovato (i costi di implementazione spesso superano il canone)",
     "https://www.capterra.it/directory/30904/cpq/software"),
]

CANALI = [
    ("Assistenza post-vendita (macchine)", "UCIMU: oltre 200 associate (più del 70% del settore, circa 8 miliardi di fatturato); parte di Federmacchine (13 associazioni, tra cui UCIMA)",
     "https://www.ucimu.it/en/association/"),
    ("Assistenza post-vendita (macchine)", "Fiere organizzate da Fieramilano: BI-MU, LAMIERA, EMO Milano (date non verificate)",
     "https://www.ucimu.it/en/association/ucimu-group/participations/"),
]

SINTESI = [
    "Assistenza post-vendita (macchine): concorrenza frammentata. Software di cataloghi ricambi e di fascicolo tecnico (italiani), field service generico "
    "a 130-230 €/utente/mese, nessuna startup italiana trovata sul tema 'assistente AI sui manuali'. Il rischio principale: i costruttori possono costruirsi un RAG "
    "da soli con strumenti generici. Il valore deve stare nel flusso di lavoro (ticket, ricambi, multilingue, storico), non nella sola chat.",
    "Incassi e solleciti: mercato molto presidiato; gli strumenti di fatturazione includono già i solleciti a 50-60 € l'anno e le piattaforme di tesoreria (Agicap, Previxio) "
    "coprono la previsione di cassa. Differenziarsi è difficile e il prezzo di riferimento è basso.",
    "Preventivatore: i CPQ per aziende grandi costano 200 €/utente/mese; per le PMI non ho trovato listini. Il rischio è di diventare consulenza (ogni cliente richiede "
    "configurazione). Un prodotto con importazione automatica di listini potrebbe evitarlo.",
    "Conclusione provvisoria: l'assistenza post-vendita per costruttori di macchine è la candidata principale. Resta da verificare con interviste che il responsabile assistenza "
    "abbia il problema e un budget.",
]

w3 = wb.create_sheet("Fase 3", 4)
w3["A1"] = "Fase 3 — concorrenti, prezzi e canali delle tre finaliste"
w3["A1"].font = H1
w3["A2"] = ("Informazioni da siti dei fornitori e stampa, quasi mai da listini ufficiali: i prezzi mancanti sono un risultato, non una dimenticanza. "
            "I punteggi nella Lista lunga sono un mio giudizio provvisorio.")
w3["A2"].font = NOTE
w3["A2"].alignment = WRAP
w3.merge_cells("A2:G2")
w3.row_dimensions[2].height = 34
hd = ["Finalista", "Concorrente / segnale", "Tipo", "Cosa offre", "Prezzo trovato", "Fonte"]
for c, (t, w) in enumerate(zip(hd, [34, 40, 26, 60, 40, 70]), start=1):
    cell = w3.cell(row=4, column=c, value=t)
    cell.font, cell.fill = H2, HEADFIL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    w3.column_dimensions[get_column_letter(c)].width = w
for i, (fin, nome, tipo, off, prezzo, url) in enumerate(COMP, start=1):
    r = 4 + i
    for c, v in enumerate([fin, nome, tipo, off, prezzo, url], start=1):
        cell = w3.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
        cell.border = THIN
        cell.font = BOLD if c == 2 else BLACK
    w3.cell(row=r, column=6).hyperlink = url
    w3.cell(row=r, column=6).font = LINK
    w3.row_dimensions[r].height = 46
r = 4 + len(COMP) + 2
w3.cell(row=r, column=1, value="Canali di vendita trovati").font = BOLD
for k, (fin, txt, url) in enumerate(CANALI, start=1):
    w3.cell(row=r + k, column=1, value=fin).font = BLACK
    c = w3.cell(row=r + k, column=2, value=txt)
    c.font, c.alignment = BLACK, WRAP
    w3.merge_cells(start_row=r + k, start_column=2, end_row=r + k, end_column=5)
    w3.cell(row=r + k, column=6, value=url).hyperlink = url
    w3.cell(row=r + k, column=6).font = LINK
    w3.row_dimensions[r + k].height = 34
r2 = r + len(CANALI) + 2
w3.cell(row=r2, column=1, value="Sintesi").font = BOLD
for k, t in enumerate(SINTESI, start=1):
    c = w3.cell(row=r2 + k, column=1, value=t)
    c.font, c.alignment = BLACK, WRAP
    w3.merge_cells(start_row=r2 + k, start_column=1, end_row=r2 + k, end_column=6)
    w3.row_dimensions[r2 + k].height = 48


wf = wb.create_sheet("Fonti")
wf.column_dimensions["A"].width = 38
wf.column_dimensions["B"].width = 120
wf["A1"] = "Fonti dei segnali (consultate fino al 30/10/2026)"
wf["A1"].font = H1
wf["A2"] = "Quasi tutte sono stampa di settore e blog di consulenza, non testi ufficiali: le date vanno confermate sulla fonte primaria."
wf["A2"].font = NOTE
fr = 4
LINK = Font(name=ARIAL, size=10, color="0563C1", underline="single")
for nome, urls in FONTI_URL.items():
    wf.cell(row=fr, column=1, value=nome).font = BOLD
    for u in urls:
        c = wf.cell(row=fr, column=2, value=u)
        c.hyperlink = u
        c.font = LINK
        fr += 1
    fr += 1

wb.save(OUT)
print("salvato:", OUT)
