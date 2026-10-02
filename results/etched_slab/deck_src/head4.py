import copy, csv, json, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn

FIG = "/tmp/deck4/fig/"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x22, 0x22, 0x22)
MUT = RGBColor(0x59, 0x59, 0x59)
TINT = RGBColor(0xEA, 0xF0, 0xF7)
BLUE, ORANGE, GREEN = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A)
GOLD = RGBColor(0xF2, 0xC1, 0x4E)
ROWLIGHT = RGBColor(0xDD, 0xE3, 0xEA)



prs = Presentation("/tmp/deck3/short.pptx")
# drop every existing slide (keep masters / layouts / logo)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
L_TITLE, L_BLANK = prs.slide_layouts[11], prs.slide_layouts[12]


def _font(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.name = "Arial"; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color


def text(slide, x, y, w, h, paras, size=14, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of str | (str, dict) | list-of-runs; dict keys: bullet, bold, size, color, level, italic, space"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple): p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        runs = p if isinstance(p, list) else [(p, {})]
        for r in runs:
            if isinstance(r, str): r = (r, {})
            rt, ro = r
            run = para.add_run(); run.text = rt
            _font(run, ro.get("size", opts.get("size", size)), ro.get("bold", opts.get("bold", False)),
                  ro.get("color", opts.get("color", color)), ro.get("italic", opts.get("italic", False)))
            if ro.get("sub"): run._r.get_or_add_rPr().set("baseline", "-25000")
        para.space_after = Pt(opts.get("space", 6))
        if opts.get("bullet"):
            lvl = opts.get("level", 0)
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.28 + 0.28 * lvl)))); pPr.set("indent", str(int(-Inches(0.22))))
            buf = etree.SubElement(pPr, qn("a:buFont")); buf.set("typeface", "Arial")
            buc = etree.SubElement(pPr, qn("a:buChar")); buc.set("char", "•" if lvl == 0 else "–")
    return tb


def title(slide, t, sub=None):
    text(slide, 0.55, 0.52, 11.2, 0.75, [(t, {"size": 30})], anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, 0.55, 1.22, 11.8, 0.45, [(sub, {"size": 15, "color": MUT, "italic": True})])


def box(slide, x, y, w, h, paras, size=13, fill=TINT):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.15 if side in ("left", "right") else 0.1))
    tf.vertical_anchor = MSO_ANCHOR.TOP
    first = True
    for p in paras:
        opts = {}
        if isinstance(p, tuple): p, opts = p
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        para.alignment = PP_ALIGN.LEFT
        runs = p if isinstance(p, list) else [(p, {})]
        for r in runs:
            if isinstance(r, str): r = (r, {})
            run = para.add_run(); run.text = r[0]
            _font(run, r[1].get("size", opts.get("size", size)), r[1].get("bold", opts.get("bold", False)),
                  r[1].get("color", opts.get("color", INK)))
            if r[1].get("sub"): run._r.get_or_add_rPr().set("baseline", "-25000")
        para.space_after = Pt(opts.get("space", 5))
    return shp


def image(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def table(slide, rows, x, y, w, col_w, size=12, row_h=0.32, bold_first_col=False, highlight=()):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for s in tblPr.findall(qn("a:tableStyleId")): tblPr.remove(s)
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            c = tbl.cell(i, j); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xFD, 0xE9, 0xD9) if i in highlight else (ROWLIGHT if i % 2 == 0 else RGBColor(0xF4, 0xF6, 0xF9)))
            tf = c.text_frame; tf.word_wrap = True
            para = tf.paragraphs[0]; para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            run = para.add_run(); run.text = str(rows[i][j])
            _font(run, size, bold=(i == 0) or (bold_first_col and j == 0), color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK)
    return gt


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def new(t, sub=None):
    s = prs.slides.add_slide(L_BLANK)
    title(s, t, sub)
    return s



import re, math
GREY = RGBColor(0x8A, 0x8A, 0x8A)
def xy_chart(slide, x, y, w, h, sers, xr, yr, xt, yt, log=False, legend=True, xmaj=1.0, ymaj=None, lsize=10):
    cd = XyChartData()
    for name, pts, *_ in sers:
        s_ = cd.add_series(name)
        for a, b in pts: s_.add_data_point(a, b)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(lsize); ch.legend.font.name = "Arial"
    va, ca = ch.value_axis, ch.category_axis
    ca.minimum_scale, ca.maximum_scale, ca.major_unit = xr[0], xr[1], xmaj
    va.minimum_scale, va.maximum_scale = yr
    if ymaj: va.major_unit = ymaj
    if log:
        sc = va._element.find(qn("c:scaling"))
        lb = etree.SubElement(sc, qn("c:logBase")); lb.set("val", "10"); sc.remove(lb); sc.insert(0, lb)
    for ax, t in ((va, yt), (ca, xt)):
        if t:
            ax.has_title = True; ax.axis_title.text_frame.text = t
            r0 = ax.axis_title.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(11); r0.font.name = "Arial"; r0.font.bold = False
        ax.tick_labels.font.size = Pt(10); ax.tick_labels.font.name = "Arial"
        ax.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
    ca.has_major_gridlines = False
    va.tick_labels.number_format_is_linked = False
    for ser, (name, pts, colr, mk, width, dash) in zip(ch.plots[0].series, sers):
        ser.format.line.color.rgb = colr; ser.format.line.width = Pt(width); ser.smooth = False
        if dash: ser.format.line.dash_style = 4
        if mk:
            ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = mk
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = colr
            ser.marker.format.line.color.rgb = colr
        else:
            ser.marker.style = XL_MARKER_STYLE.NONE
        if width == 0:
            ser.format.line.fill.background()
    return ch


def stat(slide, x, y, w, h, big, small, colr=NAVY):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08; shp.fill.solid(); shp.fill.fore_color.rgb = TINT; shp.line.fill.background(); shp.shadow.inherit = False
    text(slide, x, y + 0.12, w, h * 0.55, [(big, {"size": 30, "bold": True, "color": colr})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.1, y + h * 0.6, w - 0.2, h * 0.38, [(ln, {"size": 11, "color": MUT, "space": 0}) for ln in small.split("\n")], align=PP_ALIGN.CENTER)

