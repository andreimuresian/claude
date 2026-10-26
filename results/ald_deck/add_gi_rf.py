"""Add the group-index and RF slides to the ALD deck (AFR template), before the Conclusions.
usage: add_gi_rf.py in.pptx out.pptx"""
import sys, os, shutil, json, csv
os.makedirs("/tmp/deck5", exist_ok=True); shutil.copy(sys.argv[1], "/tmp/deck5/meeting.pptx")
sys.path.insert(0, "/home/user/claude/results/final_deck/src")
from head5 import *
from pptx.chart.data import CategoryChartData
H = "/home/user/claude/results/"
GREEN2 = RGBColor(0x1B, 0x9E, 0x77); PINK = RGBColor(0xBE, 0x18, 0x5D); RED = RGBColor(0xC0, 0x39, 0x2B)
n0 = len(prs.slides)
B = lambda t, **k: (t, dict({"size": 13, "space": 6}, **k))

# A: group index method ------------------------------------------------------------
s = new("Optical group index: how it is computed", "Same 2D FEM mode solver, three wavelengths on one fixed mesh")
box(s, 0.6, 1.8, 6.0, 4.6, [("Method", {"size": 16, "bold": True, "color": NAVY, "space": 8}),
    B("n_g = n_eff − λ · dn_eff/dλ, central difference"),
    B("Solves at λ₀ and λ₀ ± 1 nm (and ± 2 nm as a check)"),
    B("One mesh for all wavelengths: discretisation errors cancel"),
    B("All materials dispersive: LN (Zelmon), SiO₂ (Malitson), Au (Johnson & Christy), ALD Al₂O₃ (Sellmeier)", space=14),
    ("Checks", {"size": 16, "bold": True, "color": NAVY, "space": 8}),
    B("± 1 nm and ± 2 nm agree to all 5 decimals"),
    B("Light and fine meshes give the same n_g")])
box(s, 6.9, 1.8, 5.9, 4.6, [("ALD Al₂O₃ optical index", {"size": 16, "bold": True, "color": PINK, "space": 8}),
    B("No published ALD Sellmeier: fitted to ellipsometry of 20 nm ALD films (Zhukovsky et al., PRL 115, 177402, 2015)"),
    B("n² = 1 + 1.5977 λ² / (λ² − 0.010751),  λ in µm", bold=True),
    B("Error < 4·10⁻⁴ over 0.35–1.69 µm"),
    B("n = 1.614 at 1575 nm, 1.615 at 1310 nm", space=14),
    B("The earlier value n = 1.74 is closer to crystalline alumina; the Sellmeier fit is used for the conclusions", color=MUT)])

# B: group index vs ALD -------------------------------------------------------------------
s = new("Group index vs ALD thickness", "Five cross-sections; solid: ALD Sellmeier, dashed: ALD scaled to n = 1.74")
image(s, H + "ald_group_index/group_index_vs_ALD.png", 0.9, 1.7, h=4.6)
text(s, 0.6, 6.45, 12.2, 0.6, [("ALD lowers n_g by at most 0.005 (0.2 %) at 100 nm; with the buffer the drop is 2–3× smaller.",
     {"size": 14, "bold": True, "color": NAVY})])

# C: Cordoba table ----------------------------------------------------------------------------
rows = list(csv.reader(open(H + "ald_group_index/cordoba_neff_ng_table_sellmeier.tsv"), delimiter="\t"))
s = new("Cordoba at 1575 nm: rib width, buffer and ALD", "ALD Sellmeier; gap 4.15 µm, 540/270 nm, unetched slab")
table(s, [["Rib top (µm)", "SiO₂ buffer (nm)", "ALD (nm)", "n_eff", "n_g"]] + rows[1:], 0.6, 1.85, 7.3, [1.4, 1.6, 1.2, 1.55, 1.55], size=12, row_h=0.36,
      highlight=(1, 5))
box(s, 8.2, 1.85, 4.6, 4.5, [("Reading", {"size": 15, "bold": True, "color": NAVY, "space": 8}),
    B("Rib width matters most: 1.0 → 1.4 µm lowers n_g by ≈ 0.004–0.005"),
    B("Buffer alone: −0.002 to −0.003"),
    B("40 nm ALD: −0.0006 to −0.0019", space=12),
    ("Higher n_eff, lower n_g?", {"size": 14, "bold": True, "color": NAVY, "space": 6}),
    B("The wider rib confines the mode more: n_eff falls less steeply with λ (−0.227 vs −0.243 µm⁻¹), so its −λ·dn/dλ term is 0.025 smaller")])
text(s, 0.6, 5.35, 7.3, 1.0, [("All Cordoba cases lie within 2.252–2.261: < 0.5 %. ALD and buffer do not affect the optical side of velocity matching.",
     {"size": 13, "bold": True, "color": NAVY})])

# D: RF model ------------------------------------------------------------------------------------
s = new("RF study at 60 GHz: Cordoba push-pull", "Full GSG cross-section, both arms, rib 1.0 µm, cap 2.0 × 1.4 µm")
image(s, H + "rf_cordoba/rf_cross_sections.png", 0.6, 1.7, w=12.1)
box(s, 0.6, 5.1, 6.0, 1.75, [("Model", {"size": 14, "bold": True, "color": NAVY, "space": 4}),
    B("Signal 35 µm, gap 4.15 µm, grounds 49.85 µm, gold 2 µm, Si 550 µm", size=12, space=3),
    B("n_m, Z0 from the full-wave mode; R′ and L_int from the current inside the gold (volumetric, no IBC)", size=12)])
box(s, 6.8, 5.1, 6.0, 1.75, [("Layers, drawn as in the optical models", {"size": 14, "bold": True, "color": NAVY, "space": 4}),
    B("Buffer: 200 nm SiO₂ (ε 3.9) under the gold and across the gaps", size=12, space=3),
    B("ALD: 40 nm conformal Al₂O₃, ε 9.8", size=12)])

# E: RF FOMs ------------------------------------------------------------------------------------
R = [json.loads(l[7:]) for l in open(H + "rf_cordoba/cases.txt") if l.startswith("RESULT")]
def g(lp, b, a): return [r for r in R if r["lat_pad"] == lp and r["buffer_nm"] == b and r["ald_nm"] == a and r["tan_ald"] == 0 and r["mesh_factor"] == 0.35][0]
cases = [(0, 0, "Baseline"), (200, 0, "SiO₂ buffer"), (0, 40, "ALD 40 nm"), (200, 40, "Buffer + ALD")]
b0 = g(200, 0, 0)
s = new("RF figures of merit: buffer and ALD", "60 GHz, lateral padding 200 µm; changes vs baseline")
trows = [["Case", "n_m", "Z0 (Ω)", "α (dB/cm)", "Δn_m", "ΔZ0", "Δα"]]
for b, a, lab in cases:
    r = g(200, b, a)
    trows.append([lab, f"{r['n_m']:.4f}", f"{r['Z0']:.2f}", f"{r['alpha_dBcm']:.3f}"] +
                 (["–"] * 3 if b == a == 0 else [f"{(r[k] / b0[k] - 1) * 100:+.2f} %" for k in ("n_m", "Z0", "alpha_dBcm")]))
table(s, trows, 0.6, 1.8, 7.0, [1.6, 0.9, 0.9, 1.0, 0.85, 0.85, 0.9], size=12, row_h=0.38, highlight=(4,))
cd = CategoryChartData(); cd.categories = ["n_m", "Z0", "α"]
for b, a, lab in cases[1:]:
    r = g(200, b, a); cd.add_series(lab, [round((r[k] / b0[k] - 1) * 100, 2) for k in ("n_m", "Z0", "alpha_dBcm")])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.9), Inches(1.75), Inches(4.9), Inches(3.5), cd); ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(10)
for ser, c in zip(ch.plots[0].series, (BLUE, PINK, GREEN2)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c
pl = ch.plots[0]; pl.has_data_labels = True; pl.data_labels.number_format = '+0.0;-0.0'; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(9)
va = ch.value_axis; va.minimum_scale, va.maximum_scale = -5, 5; va.has_title = True; va.axis_title.text_frame.text = "change (%)"
va.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.size = Pt(11); from pptx.enum.chart import XL_TICK_LABEL_POSITION; ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW; va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
box(s, 0.6, 4.0, 7.0, 2.4, [("Why", {"size": 14, "bold": True, "color": NAVY, "space": 4}),
    B("Buffer: low-ε SiO₂ between gold and LN (ε 28–44) where the RF field is strongest → C drops: n_m ↓, Z0 ↑; R′ unchanged, so α = R′/2Z0 ↓", size=12, space=4),
    B("ALD: ε 9.8 replaces air in the gaps → C slightly up, opposite and ≈ 5× smaller", size=12, space=4),
    B("Same changes at 800 µm padding (within 0.1 %)", size=12)])
text(s, 7.9, 5.4, 4.9, 1.0, [("Buffer + ALD: n_m −2.8 %, Z0 +2.9 %, α −2.8 %", {"size": 14, "bold": True, "color": NAVY})])

# F: RF checks ----------------------------------------------------------------------------------
s = new("RF study: window, mesh and ALD loss", "Absolute values depend on the closed simulation box; the buffer / ALD changes do not")
image(s, H + "rf_cordoba/rf_window_mesh_studies.png", 0.6, 1.7, w=12.1)
stat(s, 0.6, 5.35, 3.9, 1.25, "−2.4 %", "full-wave Z0 from 100 to 1200 µm padding\n(quasi-static: ±0.03 %)", RED)
stat(s, 4.75, 5.35, 3.9, 1.25, "≤ 0.5 %", "mesh error at the chosen setting\n(factor 0.35, 126k triangles)", NAVY)
stat(s, 8.9, 5.35, 3.9, 1.25, "+0.074 dB/cm", "ALD loss at a pessimistic tan δ = 0.05\n(1.4 % of α): lossless is a fair assumption", PINK)

# G: conclusions for the new studies --------------------------------------------------------------
s = new("Conclusions: group index and RF")
items = [("Group index", "40 nm ALD ≤ 0.1 %, buffer ≈ 0.1 %, rib 1.0 → 1.4 µm 0.2 %: all negligible"),
         ("Velocity matching", "Optical side unaffected by ALD and buffer; RF n_m moves from 2.01 to 1.96 with buffer + ALD"),
         ("RF impedance and loss", "Buffer: Z0 +3.9 %, α −3.7 %; ALD: Z0 −0.6 %, α +0.6 %; together Z0 +2.9 %, α −2.8 %"),
         ("ALD at RF", "Thin film: even tan δ = 0.05 adds only 1.4 % to α"),
         ("Numerics", "Use the % changes; absolute RF values need absorbing boundaries to be window-independent")]
for i, (h, t) in enumerate(items):
    box(s, 0.9, 1.75 + i * 0.95, 11.6, 0.8, [(h + ":  ", {"size": 14, "bold": True, "color": NAVY, "space": 0})], size=14)
    text(s, 3.9, 1.83 + i * 0.95, 8.5, 0.65, [(t, {"size": 14})], anchor=MSO_ANCHOR.MIDDLE)

L = prs.slides._sldIdLst; ids = list(L); new_ids = ids[n0:]; concl = ids[n0 - 1]
for sid in new_ids[:-1]: L.remove(sid)
pos = list(L).index(concl)
for k, sid in enumerate(new_ids[:-1]): L.insert(pos + k, sid)
prs.save(sys.argv[2]); print("saved", sys.argv[2], len(prs.slides))
