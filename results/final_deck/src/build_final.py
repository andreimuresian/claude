import sys, re, csv, json, copy, itertools
sys.path.insert(0, "/tmp/deck5")
import numpy as np
from head5 import *
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI

C_BLUE, C_ORANGE, C_AQUA, C_YEL = (RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34),
                                   RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xED, 0xA1, 0x00))
GREYL, GREYD = RGBColor(0xB0, 0xB0, 0xB0), RGBColor(0x7F, 0x8C, 0x8D)
R3 = "/home/user/claude/results/three_sweeps/"
RE = "/home/user/claude/results/etched_slab/"
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

# ------------------------------------------------------------------ slide copy
USED = {str(p.partname) for p in prs.part.package.iter_parts()}
def _fresh(pn):
    m = re.match(r"^(.*?)(\d*)(\.\w+)$", pn)
    for i in itertools.count(1):
        cand = f"{m[1]}{i}{m[3]}"
        if cand not in USED:
            USED.add(cand); return cand
_ADOPTED = set()
def _adopt(part):
    if id(part) in _ADOPTED: return
    _ADOPTED.add(id(part))
    part.partname = PackURI(_fresh(str(part.partname)))
    for r in part.rels.values():
        if not r.is_external: _adopt(r.target_part)

def copy_slide(src):
    layout = [l for l in prs.slide_layouts if l.name == src.slide_layout.name][0]
    dst = prs.slides.add_slide(layout)
    for ph in list(dst.placeholders): ph._element.getparent().remove(ph._element)
    rmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE): continue
        if rel.is_external:
            rmap[rId] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True); continue
        _adopt(rel.target_part)
        rmap[rId] = dst.part.relate_to(rel.target_part, rel.reltype)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None: dst._element.cSld.insert(0, copy.deepcopy(bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")): continue
        dst.shapes._spTree.append(copy.deepcopy(el))
    for el in dst.shapes._spTree.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith(RNS) and v in rmap: el.set(k, rmap[v])
    if src.has_notes_slide and src.notes_slide.notes_text_frame.text.strip():
        notes(dst, src.notes_slide.notes_text_frame.text)
    return dst

def vline(x, y0, y1, colr=GREYL):
    return (f"_v{x}", [(x, y0), (x, y1)], colr, 0, 1.0, True)

def hide_legend_entries(ch, idx):
    leg = ch.legend._element
    for i in idx:
        le = etree.SubElement(leg, qn("c:legendEntry"))
        e = etree.SubElement(le, qn("c:idx")); e.set("val", str(i))
        d = etree.SubElement(le, qn("c:delete")); d.set("val", "1")
        leg.remove(le); leg.insert(1, le)

# ------------------------------------------------------------------ part 2: slab etch (slides 1-9 of the etch deck)
etch = Presentation("/tmp/deck5/etch.pptx")
E = list(etch.slides)
for s in E[:9]:
    copy_slide(s)

# ------------------------------------------------------------------ part 3 divider (from the etch "Appendix" divider)
div = copy_slide(E[10])
for sh in div.shapes:
    if sh.has_text_frame and "Appendix" in sh.text_frame.text:
        r = [r for p in sh.text_frame.paragraphs for r in p.runs if "Appendix" in r.text][0]
        r.text = "Final checks: robustness and cavity physics"
notes(div, "Three last sweeps: over-etch robustness of the buffered design, gap dependence of the standard lifted "
           "ripple, and the full-etch geometry with no SiO2 inside the electrodes; then why the two mirrors are equivalent.")

# ------------------------------------------------------------------ C1: buffered over-etch
rows2 = list(csv.DictReader(open(R3 + "sweep2_buffered_slab_thickness.csv")))
oe = [(275 - int(r["slab_nm"]), float(r["IL_dB_per_cm"]), float(r["VpiL_Vcm"])) for r in rows2]
oe.sort()
s = new("Buffered design: robust to over-etch",
        "200 nm SiO₂, gap 3.4 µm, x_lift 7 µm; 550 nm film, rib top 1.0 µm; slab 275 → 150 nm")
xy_chart(s, 0.4, 1.75, 4.75, 4.9, [("Vπ·L", [(a, c) for a, b, c in oe], C_BLUE, 6, 2.0, False)],
         (0, 125), (2.2, 3.0), "over-etch (nm)  [slab = 275 − over-etch]", "Vπ·L (V·cm)", legend=False, xmaj=25, ymaj=0.2)
xy_chart(s, 5.2, 1.75, 4.75, 4.9, [("IL", [(a, b) for a, b, c in oe], C_ORANGE, 6, 2.0, False)],
         (0, 125), (0.001, 0.1), "over-etch (nm)  [slab = 275 − over-etch]", "IL (dB/cm)", log=True, legend=False, xmaj=25)
v0, v15, v150 = oe[0][2], oe[1][2], oe[-1][2]
stat(s, 10.15, 1.85, 2.75, 1.5, f"+{(v15 / v0 - 1) * 100:.1f} %", "Vπ·L per 15 nm\nof over-etch", colr=C_BLUE)
stat(s, 10.15, 3.5, 2.75, 1.5, "flat", "IL vs x_lift at 150 nm\n(6.6 / 7.0 / 7.4 µm)", colr=C_AQUA)
stat(s, 10.15, 5.15, 2.75, 1.5, f"{oe[0][1] / oe[-1][1]:.0f}× lower", f"IL: {oe[0][1]:.3f} → {oe[-1][1]:.4f} dB/cm\n(275 → 150 nm slab)", colr=C_ORANGE)
notes(s, "Etch-robustness study of the buffered design (gap 3.4 um, x_lift 7 um). The film stays 550 nm and the rib top 1.0 um "
         "with 60 deg sidewalls, so deeper etch thins the slab and widens the rib base. "
         + "; ".join(f"slab {275 - a} nm: IL {b:.4f} dB/cm, VpiL {c:.3f} V*cm" for a, b, c in oe) +
         f". Vpi*L rises {(v15 / v0 - 1) * 100:.1f}% for the first 15 nm and +{(v150 / v0 - 1) * 100:.0f}% at 150 nm: thinner slab, weaker "
         "EO overlap. IL keeps falling and stays flat in x_lift (the buffer keeps the wave under the gold cut off). "
         "Optically robust; the over-etch tolerance is set by Vpi*L.")

# ------------------------------------------------------------------ C2: gap variation, standard lifted
S1 = {}
for r in csv.DictReader(open(R3 + "sweep1_legacy_gap_xlift.csv")):
    S1.setdefault(float(r["gap_um"]), []).append((float(r["x_lift_um"]), float(r["IL_dB_per_cm"]), float(r["VpiL_Vcm"])))
SUM = json.load(open(R3 + "summary.json"))
F1 = {g: SUM[f"sweep1_gap{g}"] for g in (3.9, 4.2, 5.0)}
col1 = {3.9: C_ORANGE, 4.2: GREYD, 5.0: C_BLUE}
s = new("Standard lifted: the gap sets the level, not the cavity",
        "No buffer; x_lift swept (top gap varied) at bottom gap 3.9 / 4.2 / 5.0 µm")
sers = [(f"gap {g} µm", [(x, i) for x, i, v in sorted(S1[g]) if x <= 11.05], col1[g], 3, 1.25, False) for g in (3.9, 4.2, 5.0)]
xy_chart(s, 0.4, 1.7, 6.3, 3.8, sers, (6.5, 11), (0.01, 3), "x_lift (µm)", "IL (dB/cm)", log=True, xmaj=0.5, lsize=11)
sers = [(f"gap {g} µm", [(x - g / 2, i / F1[g]["IL_inf_dB_cm"]) for x, i, v in sorted(S1[g]) if x <= 11.05], col1[g], 3, 1.25, False)
        for g in (3.9, 4.2, 5.0)]
xy_chart(s, 6.75, 1.7, 6.2, 3.8, sers, (4, 9), (0.3, 3), "cavity length L = x_lift − gap/2 (µm)", "IL / IL∞", log=True, xmaj=0.5, lsize=11)
rows = [["gap (µm)", "IL∞ (dB/cm)", "IL range (dB/cm)", "period (µm)", "2α (1/µm)", "Vπ·L (V·cm)"]]
for g in (3.9, 4.2, 5.0):
    f = F1[g]
    rows.append([f"{g}", f"{f['IL_inf_dB_cm']:.3f}", f"{f['IL_min']:.3g} – {f['IL_max']:.3g}", f"{f['period_um']:.3f}",
                 f"{f['two_alpha_per_um']:.3f}", f"{f['VpiL_Vcm']:.2f}"])
table(s, rows, 0.55, 5.65, 7.6, [1.0, 1.25, 1.65, 1.2, 1.1, 1.4], size=12, row_h=0.33)
box(s, 8.5, 5.65, 4.35, 1.32, [
    ("Same period, decay and phase vs L", {"bold": True, "size": 14, "color": NAVY}),
    ("→ the cavity is unchanged", {"size": 13}),
    ("IL∞ ∝ e^(−3.3·gap[µm]): the gap only sets how much light leaks in", {"size": 13})])
notes(s, "Standard lifted design, no buffer, 1575 nm. x_lift 6.5-11 um (0.1 um step) plus far points to 30 um; the 4.2 um "
         "gap is the earlier sweep, same code and constants. Right: IL normalised to its fitted IL_inf and plotted against the "
         "cavity length L = x_lift - gap/2 (gold inner edge to the gold/SiO2 edge): the three curves collapse, same period "
         "(0.895 um), envelope decay and phase. Only IL_inf changes: 1.10 / 0.41 / 0.030 dB/cm, i.e. exp(-3.3 um^-1 x gap), "
         "the overlap of the rib-mode tail with the gold edge. Far points (12-30 um) settle on IL_inf for every gap. "
         "Vpi*L 2.08 / 2.26 / 2.76 V*cm.")

# ------------------------------------------------------------------ C3: oxide removed
new3 = [(float(r["slab_w_um"]), float(r["IL_dB_per_cm"])) for r in csv.DictReader(open(R3 + "sweep3_full_etch_gold_at_slab_level.csv"))]
prev = {float(r["slab_w_um"]): float(r["IL_dB_per_cm"]) for r in csv.DictReader(open(RE + "slab_w_sweep.csv"))}
for r in csv.DictReader(open(RE + "followup_runs.csv")):
    if re.fullmatch(r"S_[\d.]+", r["tag"]): prev[float(r["tag"][2:])] = float(r["IL_dB_per_cm"])
F3 = SUM["sweep3_fit_slab_end_in_gold"]; CMP3 = SUM["sweep3_vs_previous_slab_in_gold"]
ww = np.linspace(5.0, 50, 300); Lw = (ww - 4.2) / 2
rho = F3["rho0"] * np.exp(-F3["two_alpha_per_um"] * Lw)
s = new("SiO₂ removed from the electrodes: same ripple",
        "Full etch, gap 4.2 µm: slab ending in the gold; previous geometry = SiO₂ lift at slab level beyond x_lift = 7 µm")
text(s, 0.5, 1.72, 4.3, 0.3, [("Previous: SiO₂ lift (gold/SiO₂ edge at 7 µm)", {"size": 12, "bold": True, "color": NAVY})])
image(s, FIG + "geo3_lift.png", 0.5, 2.02, w=4.3)
text(s, 0.5, 3.18, 4.3, 0.3, [("Now: solid gold electrodes, no SiO₂ inside", {"size": 12, "bold": True, "color": NAVY})])
image(s, FIG + "geo3_goldout.png", 0.5, 3.48, w=4.3)
stat(s, 0.5, 4.8, 2.08, 1.85, f"{CMP3['median_rel_dev'] * 100:.2f} %", "median change,\nslab ends in the gold\n(28 points)", colr=C_BLUE)
stat(s, 2.72, 4.8, 2.08, 1.85, f"{F3['IL_inf_dB_cm']:.3f}", "dB/cm asymptote\n(Cordoba 0.42)", colr=C_ORANGE)
# longest series first: LibreOffice drops shorter XY series listed before a longer one
sers = [("_envU", list(zip(ww, F3["IL_inf_dB_cm"] * (1 + rho) / (1 - rho))), RGBColor(0xF3, 0xB0, 0x94), 0, 0.75, True),
        ("_envL", list(zip(ww, F3["IL_inf_dB_cm"] * (1 - rho) / (1 + rho))), RGBColor(0xF3, 0xB0, 0x94), 0, 0.75, True),
        ("previous (SiO₂ lift)", [p for p in sorted(prev.items()) if p[1] >= 0.001], GREYD, 0, 1.25, False),
        ("SiO₂ removed", [p for p in sorted(new3) if p[1] >= 0.001], C_ORANGE, 3, 1.0, False),
        (f"IL∞ = {F3['IL_inf_dB_cm']:.3f} dB/cm", [(4.2, F3["IL_inf_dB_cm"]), (50, F3["IL_inf_dB_cm"])], C_ORANGE, 0, 1.0, True),
        vline(14.0, 0.001, 10, C_BLUE)]
ch = xy_chart(s, 5.0, 1.7, 7.95, 5.0, sers, (2, 50), (0.001, 10), "slab_w (µm)", "IL (dB/cm)", log=True, xmaj=4, lsize=11)
hide_legend_entries(ch, [0, 1, 5])
text(s, 9.15, 5.42, 2.4, 0.3, [("old gold/SiO₂ edge (slab_w 14)", {"size": 10, "color": C_BLUE})])
notes(s, "Full etch, gap 4.2 um, x_lift 7 um. Coworker's test: remove all SiO2 from the electrodes (lower block, column and pad "
         "fused in one solid gold body out to 32 um, air outside), so the slab can only end inside the gold. "
         f"Slab ending between the gap and 14 um: identical to the previous sweep (median {CMP3['median_rel_dev'] * 100:.2f} %, max "
         f"{CMP3['max_rel_dev'] * 100:.1f} % at slab_w 13.95, next to where the SiO2 used to start). Past 14 um the flat 0.228 dB/cm "
         f"line is gone: the slab end remains the only mirror and the ripple decays slowly to IL_inf = {F3['IL_inf_dB_cm']:.3f} dB/cm "
         f"(Airy fit: period {F3['period_um']:.3f} um of slab-end position, 2alpha {F3['two_alpha_per_um']:.3f}/um, rho0 {F3['rho0']:.3f}), "
         "the Cordoba (no reflection) level. Dotted: Airy envelope. Above 30 um the 1.6 um step aliases the 1.79 um ripple period. "
         "A full etch ending in the gold therefore makes IL depend on the etch-edge position with a 0.9 um period, consistent with "
         "the messy IL measured on the earlier full-etch tapeout.")

# ------------------------------------------------------------------ C4: why the two mirrors are the same
fitc = json.load(open(RE + "ripple_comparison_facet_vs_interface.json"))
A_, B_ = fitc["slab_facet_in_gold"], fitc["gold_SiO2_interface"]
s = new("Why the slab end and the gold edge act the same",
        "The lateral wave lives only where gold and LN slab overlap (1D stack indices, 1575 nm)")
image(s, FIG + "mech_two_mirrors.png", 0.35, 1.7, w=8.75)
box(s, 9.3, 1.75, 3.6, 2.75, [
    ("Gold + LN: n = 2.07 > 1.884", {"bold": True, "size": 14, "color": NAVY}),
    ("→ propagates sideways", {"size": 13, "space": 8}),
    ("Remove the LN: 1.46", {"bold": True, "size": 14, "color": NAVY}),
    ("Remove the gold: 1.61", {"bold": True, "size": 14, "color": NAVY}),
    ("→ both evanescent → reflection", {"size": 13, "space": 8}),
    ("Mirror = end of the overlap, whatever ends it", {"bold": True, "size": 13.5})])
rows = [["fit", "slab end", "gold edge"],
        ["period (µm)", f"{A_['period_um']:.3f}", f"{B_['period_um']:.3f}"],
        ["IL∞ (dB/cm)", f"{A_['IL_inf']:.3f}", f"{B_['IL_inf']:.3f}"],
        ["decay 2α (1/µm)", f"{A_['two_alpha_per_um']:.3f}", f"{B_['two_alpha_per_um']:.3f}"],
        ["mirror ρ₀", f"{A_['rho0']:.3f}", f"{B_['rho0']:.3f}"]]
table(s, rows, 9.3, 4.7, 3.6, [1.5, 1.05, 1.05], size=12, row_h=0.33)
text(s, 9.3, 6.4, 3.6, 0.6, [("Cavity sets period, decay, IL∞; the mirror only ρ₀ and phase", {"size": 11.5, "color": MUT, "italic": True})])
notes(s, "The lateral wave is a hybrid (plasmon-assisted) mode of the vertical stack gold / LN slab / BOX. It needs both: the metal pulls "
         "the field to the top of the slab and raises its index to 2.07, above the rib mode (1.884), so it can be phase matched "
         "and propagate sideways (kx real, lateral wavelength 1.79 um, ripple period 0.9 um). "
         "Remove the LN (slab end inside the gold): the stack becomes gold on BOX, a gold/SiO2 surface plasmon with index 1.46. "
         "Remove the gold (gold edge, slab continues under SiO2): the stack becomes an oxide-clad LN slab, index 1.61. "
         "Both are below 1.884: kx becomes imaginary, the wave cannot continue and is reflected. "
         "So the mirror is simply the end of the gold-LN overlap; its height, the material that replaces the stack and the "
         "length of the overlap do not matter for the period, the decay and IL_inf, which belong to the overlap region. "
         "The termination only sets the reflection strength and phase: the gold-wrapped slab end is a hard mirror (node at the wall, "
         "rho0 0.957), the gold edge a soft one (evanescent tail into the oxide-clad slab, rho0 0.922; effective planes 0.3 um apart, "
         "hence the 0.34 um shift between the two IL curves). Indices: 1D TM model of each stack, average of the two in-plane "
         "LN permittivities, Sellmeier SiO2 and Johnson & Christy gold.")

# ------------------------------------------------------------------ conclusions
f34 = {r["slab_nm"]: r for r in rows2}
s = new("Conclusions", "Lifted electrodes, λ = 1575 nm")
stat(s, 0.6, 1.8, 3.9, 1.65, f"{float(f34['275']['IL_dB_per_cm']):.3f} dB/cm", "IL, flat vs x_lift\n200 nm buffer, gap 3.4 µm", colr=C_BLUE)
stat(s, 4.72, 1.8, 3.9, 1.65, f"{float(f34['275']['VpiL_Vcm']):.2f} V·cm", "Vπ·L at gap 3.4 µm\n(+2.6 % per 15 nm over-etch)", colr=C_BLUE)
stat(s, 8.84, 1.8, 3.9, 1.65, "38 nm", "buffer cutoff\n200 nm = 5× margin", colr=C_BLUE)
box(s, 0.6, 3.75, 12.14, 2.6, [
    ("Cause: the wave under the gold exists only where gold and LN slab overlap; any end of the overlap is a mirror → lateral cavity, period 0.9 µm", {"size": 15, "space": 9}),
    ("Gap: scales the leak (IL∞ ∝ e^(−3.3·gap)), not the cavity", {"size": 15, "space": 9}),
    ("Slab etching: the slab end in the gold is just another mirror; ending it in the gap costs +35 … +140 % Vπ·L; thin LN under the gold costs 2 – 3× IL at equal Vπ·L", {"size": 15, "space": 9}),
    ("Buffer ≥ 38 nm SiO₂ removes the wave: keep 200 nm (no new material, no extra litho, robust to over-etch)", {"size": 15, "bold": True, "color": NAVY, "space": 9}),
])
notes(s, "Final recommendation: 200 nm SiO2 buffer, bottom gap 3.4 um: IL 0.026 dB/cm flat in x_lift, Vpi*L 2.27 V*cm. "
         "Over-etch: IL keeps decreasing and stays flat; Vpi*L +2.6% per 15 nm (+27% at 150 nm slab). "
         "The full-etch geometry ending in the gold reproduces the ripple whatever surrounds the electrodes; the messy IL of the "
         "earlier full-etch tapeout is consistent with this.")

out = "/tmp/deck5/VPI_Lisbon_Lifted_final.pptx"
prs.save(out); print("saved", out, len(prs.slides))
